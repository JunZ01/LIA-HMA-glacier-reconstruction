from __future__ import annotations

import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


EXPECTED_SHA256 = "8ffc24620f1189d4ed49e441dda4332552a93ec744abbd1cf0e13823d55462c8"
KEYWORDS = ("chronolog", "reconstruct", "temperature", "table", "figure", "fig.", "fig ", "data")


def project_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / "LIA_PROJECT_FACTS.md").exists():
            return parent
    raise RuntimeError("project-root-not-found")


def fingerprint(path: Path) -> dict[str, object]:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    stat = path.stat()
    return {"path": str(path), "size_bytes": stat.st_size, "mtime_ns": stat.st_mtime_ns, "sha256": digest}


def main() -> int:
    root = project_root()
    loop = root / "validation_data/results/paleo_surface_validation_v0/climate_reconstruction_improvement_loop_v2"
    output = loop / "round_248_npl065_supplement_structure_audit_v1"
    output.mkdir(parents=True, exist_ok=False)
    source = loop / (
        "round_247_npl065_official_supplement_acquisition_v2/raw_unopened/"
        "1-s2.0-S1125786520301387-mmc1.docx"
    )
    contract = root / (
        "Himalaya_greatGlacier/With_Rec_climate_organized/ice_surface_reconstruction/drivers/"
        "round248_npl065_supplement_structure_audit_v1_contract.md"
    )
    before = {"source": fingerprint(source), "contract": fingerprint(contract), "driver": fingerprint(Path(__file__).resolve())}
    if before["source"]["sha256"] != EXPECTED_SHA256:
        raise RuntimeError("supplement-hash-changed")

    document = Document(source)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append({"table_index": table_index, "row_count": len(rows), "column_count": max((len(row) for row in rows), default=0), "rows": rows})
    with zipfile.ZipFile(source, "r") as archive:
        names = sorted(archive.namelist())
        media = [name for name in names if name.startswith("word/media/")]
        embeddings = [name for name in names if name.startswith("word/embeddings/")]
        charts = [name for name in names if name.startswith("word/charts/")]
    keyword_paragraphs = [
        {"paragraph_index": index, "text": text}
        for index, text in enumerate(paragraphs, start=1)
        if any(keyword in text.lower() for keyword in KEYWORDS)
    ]
    captions = [
        {"paragraph_index": index, "text": text}
        for index, text in enumerate(paragraphs, start=1)
        if re.match(r"^(fig(?:ure)?|table)\s*\.?\s*[sS]?\d+", text, flags=re.IGNORECASE)
    ]
    machine_readable_annual_table_candidates = []
    for table in tables:
        year_like_rows = 0
        for row in table["rows"]:
            if row and re.fullmatch(r"(?:17|18|19|20)\d{2}", row[0].strip()):
                year_like_rows += 1
        if year_like_rows >= 10:
            machine_readable_annual_table_candidates.append({"table_index": table["table_index"], "year_like_row_count": year_like_rows})

    text_inventory = []
    for index, text in enumerate(paragraphs, start=1):
        text_inventory.append(f"P{index:04d}\t{text}")
    for table in tables:
        text_inventory.append(f"TABLE {table['table_index']} {table['row_count']}x{table['column_count']}")
        for row_index, row in enumerate(table["rows"], start=1):
            text_inventory.append(f"T{table['table_index']:03d}R{row_index:04d}\t" + "\t".join(row))
    (output / "text_and_table_inventory.txt").write_text("\n".join(text_inventory) + "\n", encoding="utf-8")

    after = {"source": fingerprint(source), "contract": fingerprint(contract), "driver": fingerprint(Path(__file__).resolve())}
    unchanged = before == after
    if not unchanged:
        raise RuntimeError("source-fingerprint-changed")
    has_numerical_series = bool(machine_readable_annual_table_candidates or embeddings or charts)
    decision = (
        "inspect_machine_readable_supplement_series"
        if has_numerical_series
        else "official_supplement_has_no_machine_readable_chronology_or_reconstruction"
    )
    result = {
        "experiment_id": output.name,
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "status": "complete",
        "source": before["source"],
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "table_shapes": [[table["row_count"], table["column_count"]] for table in tables],
        "inline_shape_count": len(document.inline_shapes),
        "media_part_count": len(media),
        "media_parts": media,
        "embedding_part_count": len(embeddings),
        "embedding_parts": embeddings,
        "chart_part_count": len(charts),
        "chart_parts": charts,
        "keyword_paragraphs": keyword_paragraphs,
        "captions": captions,
        "machine_readable_annual_table_candidates": machine_readable_annual_table_candidates,
        "has_machine_readable_numerical_series_candidate": has_numerical_series,
        "decision": decision,
        "embedded_images_extracted_or_rendered": False,
        "annual_proxy_or_climate_values_decoded": False,
        "locked_1921_1939_values_accessed": False,
        "jionglacuo_opened_or_hashed": False,
        "binary_executed": False,
        "enkf_run": False,
        "production_reconstruction_run": False,
        "oggm_run": False,
        "source_fingerprints_unchanged": unchanged,
    }
    (output / "summary.json").write_text(json.dumps(result, indent=2), encoding="ascii")
    (output / "source_manifest.json").write_text(json.dumps({"before": before, "after": after}, indent=2), encoding="ascii")
    (output / "FINAL_RESULT.md").write_text(
        "# Round 248 NPL065 supplement structure audit\n\n"
        f"- Paragraphs: `{len(paragraphs)}`; tables: `{len(tables)}`; inline shapes: `{len(document.inline_shapes)}`.\n"
        f"- Embedded files: `{len(embeddings)}`; chart parts: `{len(charts)}`; media parts: `{len(media)}`.\n"
        f"- Machine-readable annual-series candidates: `{len(machine_readable_annual_table_candidates)}`.\n"
        f"- Images extracted or rendered: `false`.\n\n"
        f"Decision: `{decision}`.\n",
        encoding="ascii",
    )
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
