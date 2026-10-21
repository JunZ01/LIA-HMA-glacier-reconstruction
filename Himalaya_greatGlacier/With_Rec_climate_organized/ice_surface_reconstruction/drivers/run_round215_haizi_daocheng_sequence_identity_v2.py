from __future__ import annotations

import hashlib
import json
import math
import urllib.request
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
LOOP = (
    ROOT
    / "validation_data/results/paleo_surface_validation_v0"
    / "climate_reconstruction_improvement_loop_v2"
)
OUT = LOOP / "round_215_haizi_daocheng_sequence_identity_v1"
WORKBOOK = (
    LOOP
    / "round_214_haizi_annual_temperature_data_recovery_v1"
    / "41598_2024_79096_MOESM1_ESM.xlsx"
)
CONTRACT = OUT / "ROUND_CONTRACT.md"
DRIVER = Path(__file__).resolve()
NOAA_URL = (
    "https://www.ncei.noaa.gov/pub/data/paleo/pages2k/"
    "pages2k-temperature-v2-2017/data-version-2.0.0/"
    "Asi-CHIN018.Li.2010.txt"
)
EXPECTED_WORKBOOK_SHA256 = (
    "35302c0ba1912b015ee8d8b1a5c05df9aa82dd44d4467f7b07278e723611d3eb"
)
LOCKED_START = 1921
LOCKED_END = 1939


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def state(path: Path) -> dict[str, object]:
    return {"path": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)}


def download(url: str, destination: Path) -> None:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 LIA climate evidence audit"},
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        destination.write_bytes(response.read())


def read_haizi_year_first() -> tuple[dict[int, float], int]:
    if sha256(WORKBOOK) != EXPECTED_WORKBOOK_SHA256:
        raise RuntimeError("Haizi workbook hash changed")
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    if [sheet.cell(1, column).value for column in range(1, 5)] != [
        "year",
        "std",
        "Temperature(��)",
        "num",
    ]:
        raise RuntimeError("Unexpected Haizi workbook header")
    values: dict[int, float] = {}
    skipped = 0
    for row in sheet.iter_rows(min_row=2, max_col=4, values_only=False):
        year_token = row[0].value
        if year_token is None:
            continue
        year = int(year_token)
        if year < 1115 or year > 2017:
            continue
        if LOCKED_START <= year <= LOCKED_END:
            skipped += 1
            continue
        value = float(row[1].value)
        if not math.isfinite(value):
            raise RuntimeError(f"Non-finite Haizi STD in allowed year {year}")
        values[year] = value
    workbook.close()
    expected = set(range(1115, 1921)) | set(range(1940, 2018))
    if set(values) != expected or skipped != 19:
        raise RuntimeError("Haizi allowed coverage mismatch")
    return values, skipped


def read_noaa_year_first(path: Path) -> tuple[dict[int, float], int]:
    values: dict[int, float] = {}
    skipped = 0
    data_started = False
    with path.open("rb") as stream:
        for line_number, raw in enumerate(stream, start=1):
            stripped = raw.strip()
            if not stripped or stripped.startswith(b"#"):
                continue
            if not data_started:
                if stripped.lower().startswith(b"year"):
                    data_started = True
                    continue
                raise RuntimeError(f"Unexpected NOAA line before header: {line_number}")
            separator = stripped.find(b"\t")
            if separator <= 0:
                raise RuntimeError(f"Missing NOAA year separator: {line_number}")
            year = int(float(stripped[:separator].decode("ascii")))
            if LOCKED_START <= year <= LOCKED_END:
                skipped += 1
                continue
            remainder = stripped[separator + 1 :].split()[0]
            value = float(remainder.decode("ascii"))
            if not math.isfinite(value):
                raise RuntimeError(f"Non-finite NOAA value in allowed year {year}")
            values[year] = value
    expected = set(range(1540, 1921)) | set(range(1940, 2007))
    if set(values) != expected or skipped != 19:
        raise RuntimeError("NOAA Daocheng allowed coverage mismatch")
    return values, skipped


def correlation(a: np.ndarray, b: np.ndarray) -> float:
    return float(np.corrcoef(a, b)[0, 1])


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    noaa = OUT / "Asi-CHIN018.Li.2010.txt"
    if not noaa.exists():
        download(NOAA_URL, noaa)
    sources = {"workbook": WORKBOOK, "noaa": noaa, "contract": CONTRACT, "driver": DRIVER}
    for path in sources.values():
        if not path.exists():
            raise FileNotFoundError(path)
    before = {name: state(path) for name, path in sources.items()}

    haizi, haizi_skipped = read_haizi_year_first()
    daosic, daosic_skipped = read_noaa_year_first(noaa)
    common_years = sorted(set(haizi) & set(daosic))
    if len(common_years) != 448:
        raise RuntimeError(f"Unexpected allowed overlap count: {len(common_years)}")
    x = np.asarray([daosic[year] for year in common_years], dtype=float)
    y = np.asarray([haizi[year] for year in common_years], dtype=float)
    raw_r = correlation(x, y)
    d1_r = correlation(np.diff(x), np.diff(y))
    x_smooth = pd.Series(x).rolling(31, center=True, min_periods=31).mean().to_numpy()
    y_smooth = pd.Series(y).rolling(31, center=True, min_periods=31).mean().to_numpy()
    valid_smooth = np.isfinite(x_smooth) & np.isfinite(y_smooth)
    smooth_r = correlation(x_smooth[valid_smooth], y_smooth[valid_smooth])
    slope, intercept = np.polyfit(x, y, 1)
    residual = y - (intercept + slope * x)
    residual_rmse = float(np.sqrt(np.mean(residual**2)))
    normalized_rmse = residual_rmse / float(np.std(y, ddof=1))

    if abs(raw_r) >= 0.9:
        classification = "shared_or_derived_chronology_material"
        admissible = False
    elif abs(raw_r) >= 0.7:
        classification = "related_chronology_material_not_independent"
        admissible = False
    else:
        classification = "sequence_independent_for_annual_temperature_gate"
        admissible = True

    rows = pd.DataFrame(
        {
            "year": common_years,
            "daosic_pages2k_trsgi": x,
            "haizi_2024_std": y,
            "daosic_z": (x - np.mean(x)) / np.std(x, ddof=1),
            "haizi_z": (y - np.mean(y)) / np.std(y, ddof=1),
            "daosic_smooth31": x_smooth,
            "haizi_smooth31": y_smooth,
        }
    )
    rows.to_csv(OUT / "allowed_sequence_comparison.csv", index=False)

    figure, axes = plt.subplots(2, 1, figsize=(12.4, 7.4), sharex=True)
    axes[0].plot(common_years, rows.daosic_z, color="#B33A3A", linewidth=1.0, label="PAGES2k Daocheng")
    axes[0].plot(common_years, rows.haizi_z, color="#175B8C", linewidth=1.0, label="Haizi 2024 STD")
    axes[0].axvspan(1921, 1939, color="#E5E5E5", alpha=0.85, linewidth=0)
    axes[0].set_ylabel("Standardized annual index")
    axes[0].set_title("Haizi 2024 versus current Daocheng PAGES2k chronology")
    axes[0].grid(axis="y", color="#DDDDDD", linewidth=0.6)
    axes[0].legend(frameon=False, ncol=2, loc="upper left")
    axes[1].plot(common_years, x_smooth, color="#B33A3A", linewidth=2.0, label="PAGES2k Daocheng")
    axes[1].plot(common_years, y_smooth, color="#175B8C", linewidth=2.0, label="Haizi 2024 STD")
    axes[1].axvspan(1921, 1939, color="#E5E5E5", alpha=0.85, linewidth=0)
    axes[1].set_xlabel("Year")
    axes[1].set_ylabel("31-year mean index")
    axes[1].grid(axis="y", color="#DDDDDD", linewidth=0.6)
    axes[1].legend(frameon=False, ncol=2, loc="upper left")
    figure.subplots_adjust(left=0.09, right=0.98, top=0.93, bottom=0.09, hspace=0.16)
    figure.savefig(OUT / "haizi_daocheng_sequence_identity.png", dpi=220)
    plt.close(figure)

    after = {name: state(path) for name, path in sources.items()}
    if before != after:
        raise RuntimeError("A read-only source changed")
    summary = {
        "round": 215,
        "decision": classification,
        "admissible_as_independent_annual_temperature_evidence": admissible,
        "overlap": {
            "allowed_common_year_count": len(common_years),
            "first_year": common_years[0],
            "last_year": common_years[-1],
        },
        "metrics": {
            "raw_correlation": raw_r,
            "first_difference_correlation": d1_r,
            "smooth31_correlation": smooth_r,
            "affine_slope_haizi_on_daosic": float(slope),
            "affine_intercept_haizi_on_daosic": float(intercept),
            "affine_residual_rmse": residual_rmse,
            "affine_residual_rmse_fraction_of_haizi_sd": normalized_rmse,
        },
        "locked_parser": {
            "haizi_locked_rows_skipped_before_std_access": haizi_skipped,
            "daosic_locked_rows_skipped_before_trsgi_access": daosic_skipped,
        },
        "source_states_unchanged": True,
        "locked_period_values_accessed": False,
        "jionglacuo_accessed": False,
    }
    (OUT / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=True) + "\n", encoding="utf-8"
    )
    (OUT / "source_manifest.json").write_text(
        json.dumps(
            {"sources_before": before, "sources_after": after, "noaa_url": NOAA_URL},
            indent=2,
            ensure_ascii=True,
        )
        + "\n",
        encoding="utf-8",
    )
    result = f"""# Round 215 Result

`{classification}`

Across {len(common_years)} allowed common years, raw, first-difference, and
31-year correlations are `{raw_r:.4f}`, `{d1_r:.4f}`, and `{smooth_r:.4f}`.
The affine residual RMSE is `{normalized_rmse:.3f}` of the Haizi chronology
standard deviation.

Independent annual-temperature admission: `{str(admissible).lower()}`.
No locked-period or Jionglacuo value was accessed.
"""
    (OUT / "FINAL_RESULT.md").write_text(result, encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
