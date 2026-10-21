"""Blind western-HMA annual-temperature validation against Chongce Core 3."""

from __future__ import annotations

import json
from pathlib import Path

import h5py
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

import validate_temperature_external_low_frequency_v1 as base


SOURCE = (
    base.ROOT
    / "validation_data/sources/external_hma_temperature_candidates_lake_ice_v1"
    / "chongce_zenodo_5766235.xlsx"
)
PUBLIC = Path(
    r"E:\3spyderProject\HMA_500\Down_scale_new\out_data\climate_anomaly_05deg_with_ensemble_1000-2000.nc"
)
LEDGER = (
    base.RESULTS
    / "climate_reconstruction_improvement_loop_v2/round_107_proxy_evidence_master_ledger_v1"
    / "proxy_evidence_master_ledger.csv"
)
OUTPUT = (
    base.RESULTS
    / "climate_reconstruction_temperature_chongce_independent_blind_validation_v1"
)
VARIABLE = "tas_ensemble_mean_rec_result"
YEARS = np.arange(1000, 2001)
WINDOWS = [(year, year + 49) for year in range(1550, 1900, 50)]


def read_chongce(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = []
    for row in sheet.iter_rows(min_row=6):
        try:
            year = int(float(row[0].value))
        except (TypeError, ValueError):
            continue
        if base.LOCKED_START <= year <= base.LOCKED_END:
            continue
        value = row[2].value
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        rows.append({"year": year, "external": float(value)})
    workbook.close()
    return pd.DataFrame(rows).sort_values("year").reset_index(drop=True)


def field_without_locked(dataset: h5py.Dataset) -> np.ndarray:
    result = np.full(dataset.shape, np.nan, dtype=float)
    result[:921] = np.asarray(dataset[:921], dtype=float)
    result[940:] = np.asarray(dataset[940:], dtype=float)
    return result


def local_series(field: np.ndarray, lat: np.ndarray, lon: np.ndarray) -> tuple[np.ndarray, int]:
    mask = (
        (lat[:, None] >= 34.75)
        & (lat[:, None] <= 35.75)
        & (lon[None, :] >= 80.75)
        & (lon[None, :] <= 81.75)
    )
    weights = np.cos(np.deg2rad(lat))[:, None] * mask
    values = np.where(mask[None, :, :], field, np.nan)
    valid = np.isfinite(values)
    numerator = np.nansum(values * weights[None, :, :], axis=(1, 2))
    denominator = np.sum(valid * weights[None, :, :], axis=(1, 2))
    return (
        np.divide(numerator, denominator, out=np.full(len(field), np.nan), where=denominator > 0),
        int(mask.sum()),
    )


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contract = OUTPUT / "ROUND_CONTRACT.md"
    if not contract.exists():
        raise FileNotFoundError("Contract must exist before numeric parsing")
    paths = {
        "chongce": SOURCE,
        "public_field": PUBLIC,
        "frozen_candidate": base.CANDIDATE,
        "active_proxy_ledger": LEDGER,
        "contract": contract,
    }
    before = {name: base.sha256(path) for name, path in paths.items()}

    ledger = pd.read_csv(LEDGER)
    nearby = ledger[
        ledger.latitude.between(34.2349, 36.2349)
        & ledger.longitude.between(80.1144, 82.1144)
    ]
    if len(nearby) != 0:
        raise RuntimeError("Chongce is not spatially held out from active proxy ledger")

    external = read_chongce(SOURCE)
    with h5py.File(PUBLIC, "r") as handle:
        years = np.asarray(handle["time"], dtype=int).ravel()
        lat = np.asarray(handle["lat"], dtype=float).ravel()
        lon = np.asarray(handle["lon"], dtype=float).ravel()
        if not np.array_equal(years, YEARS):
            raise ValueError("Unexpected field years")
        project, cell_count = local_series(
            field_without_locked(handle[VARIABLE]), lat, lon
        )

    local = pd.DataFrame({"year": YEARS, "project": project})
    frozen = base.read_candidate(base.CANDIDATE)
    frozen["increment"] = frozen.candidate - frozen.project
    local = local.merge(frozen[["year", "increment"]], on="year", how="left")
    local["candidate"] = local.project + local.increment
    merged = local.merge(external, on="year", how="inner")

    modern = merged.year.between(1940, 2000)
    if modern.sum() != 61:
        raise ValueError("Modern alignment coverage is not 61 years")
    for column in ["project", "candidate", "external"]:
        merged[f"{column}_aligned"] = merged[column] - merged.loc[modern, column].mean()

    rows = []
    for start, end in WINDOWS:
        subset = merged.loc[merged.year.between(start, end)]
        if len(subset) != 50:
            raise ValueError(f"Incomplete Chongce window {start}-{end}")
        rows.append(
            {
                "window_start": start,
                "window_end": end,
                "project_degC": subset.project_aligned.mean(),
                "candidate_degC": subset.candidate_aligned.mean(),
                "chongce_degC": subset.external_aligned.mean(),
            }
        )
    windows = pd.DataFrame(rows)
    project_error = windows.project_degC - windows.chongce_degC
    candidate_error = windows.candidate_degC - windows.chongce_degC
    project_rmse = float(np.sqrt(np.mean(project_error**2)))
    candidate_rmse = float(np.sqrt(np.mean(candidate_error**2)))
    improvement = (project_rmse - candidate_rmse) / project_rmse
    project_bias = float(project_error.mean())
    candidate_bias = float(candidate_error.mean())
    project_corr = float(windows[["project_degC", "chongce_degC"]].corr().iloc[0, 1])
    candidate_corr = float(windows[["candidate_degC", "chongce_degC"]].corr().iloc[0, 1])
    gates = {
        "rmse_improves_at_least_10pct": bool(improvement >= 0.10 - 1e-12),
        "absolute_mean_bias_decreases": bool(abs(candidate_bias) < abs(project_bias)),
        "window_correlation_decline_at_most_0p10": bool(candidate_corr - project_corr >= -0.10 - 1e-12),
    }
    passes = bool(all(gates.values()))

    windows.to_csv(OUTPUT / "chongce_50year_windows.csv", index=False)
    pd.DataFrame(
        [
            {
                "project_rmse_degC": project_rmse,
                "candidate_rmse_degC": candidate_rmse,
                "rmse_improvement_fraction": improvement,
                "project_mean_bias_degC": project_bias,
                "candidate_mean_bias_degC": candidate_bias,
                "project_window_correlation": project_corr,
                "candidate_window_correlation": candidate_corr,
                "grid_cell_count": cell_count,
            }
        ]
    ).to_csv(OUTPUT / "chongce_metrics.csv", index=False)

    x = (windows.window_start + windows.window_end) / 2
    fig, axis = plt.subplots(figsize=(10.5, 5.2), constrained_layout=True)
    axis.plot(x, windows.chongce_degC, color="#111111", marker="o", label="Chongce annual")
    axis.plot(x, windows.project_degC, color="#d55e00", marker="s", label="Project local")
    axis.plot(x, windows.candidate_degC, color="#0072b2", marker="^", label="Frozen increment")
    axis.axhline(0, color="#888888", linewidth=0.8)
    axis.set_title("Western HMA annual temperature: Chongce blind validation")
    axis.set_xlabel("50-year window midpoint")
    axis.set_ylabel("Temperature anomaly relative to 1940-2000 (degC)")
    axis.grid(alpha=0.2)
    axis.legend()
    fig.savefig(OUTPUT / "chongce_blind_validation.png", dpi=220)
    plt.close(fig)

    after = {name: base.sha256(path) for name, path in paths.items()}
    (OUTPUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "inputs": {
                    name: {"path": str(paths[name]), "sha256_before": before[name], "sha256_after": after[name]}
                    for name in paths
                },
                "all_input_hashes_unchanged": before == after,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    result = {
        "status": "complete",
        "spatially_held_out_from_active_proxy_ledger": True,
        "project_rmse_degC": project_rmse,
        "candidate_rmse_degC": candidate_rmse,
        "rmse_improvement_fraction": improvement,
        "project_mean_bias_degC": project_bias,
        "candidate_mean_bias_degC": candidate_bias,
        "project_window_correlation": project_corr,
        "candidate_window_correlation": candidate_corr,
        "gates": gates,
        "passes_chongce_gates": passes,
        "locked_1921_1939_workbook_values_accessed": False,
        "locked_1921_1939_field_slices_read": False,
        "phase_1901_1920_used": False,
        "spatial_field_release_authorized": False,
        "production_release_authorized": False,
        "decision": (
            "western_annual_evidence_supports_frozen_increment_diagnostic"
            if passes
            else "western_annual_evidence_rejects_frozen_increment_diagnostic"
        ),
    }
    (OUTPUT / "summary.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    (OUTPUT / "FINAL_RESULT.md").write_text(
        "# Chongce Independent Blind Validation\n\n"
        f"- Project RMSE: `{project_rmse:.3f} degC`; candidate RMSE: `{candidate_rmse:.3f} degC`.\n"
        f"- RMSE improvement: `{100*improvement:.2f}%`.\n"
        f"- Mean bias: `{project_bias:+.3f}` to `{candidate_bias:+.3f} degC`.\n"
        f"- All gates pass: `{str(passes).lower()}`.\n"
        "- Western diagnostic only; no spatial/production field is released.\n",
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
