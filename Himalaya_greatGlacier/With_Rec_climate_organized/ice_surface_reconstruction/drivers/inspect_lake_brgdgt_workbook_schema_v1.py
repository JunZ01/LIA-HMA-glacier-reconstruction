from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
EXPERIMENT = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "temperature_lake_brgdgt_recalibration_sensitivity_v1"
)
SOURCE = EXPERIMENT / "source/Liang_et_al_20241108_CEE_Supplementary_Data.xlsx"
OUTPUT = EXPERIMENT / "workbook_schema.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        string_cells = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25)):
            for cell in row:
                # Numeric cells can contain ages or temperatures. The schema pass
                # deliberately never decodes them.
                if cell.data_type in {"s", "inlineStr"} and cell.value is not None:
                    string_cells.append({"coordinate": cell.coordinate, "text": str(cell.value)})
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "string_cells_first_25_rows": string_cells,
            }
        )
    workbook.close()
    result = {
        "experiment": "temperature_lake_brgdgt_recalibration_sensitivity_v1",
        "source": str(SOURCE),
        "source_sha256": sha256(SOURCE),
        "source_size_bytes": SOURCE.stat().st_size,
        "sheets": sheets,
        "numeric_cells_decoded": False,
        "sealed_period_values_used": False,
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    OUTPUT.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
