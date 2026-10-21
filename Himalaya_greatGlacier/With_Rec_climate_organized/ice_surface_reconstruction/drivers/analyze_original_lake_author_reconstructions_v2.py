from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from compare_noaa_hma_lake_temperature_project_v1 import (
    LINEAGE,
    PROJECT_MEMBERS,
    read_project_series,
)


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
BASE = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2"
)
SOURCE_ROOT = BASE / "temperature_original_lake_author_supplements_v1/source"
LIANG_SOURCE = (
    BASE
    / "temperature_lake_brgdgt_recalibration_sensitivity_v1/source/"
    "Liang_et_al_20241108_CEE_Supplementary_Data.xlsx"
)
OUTPUT = BASE / "temperature_original_lake_author_reconstructions_v2"
DRIVER = Path(__file__).resolve()

LOCKED_START = 1921.0
LOCKED_END = 1939.999999
PRELOCK_START = 1500.0
PRELOCK_END = 1920.999999
POSTLOCK_START = 1940.0
POSTLOCK_END = 1950.0

RECORDS = {
    "Tengchongqinghai": {
        "source": SOURCE_ROOT / "zhao2021_lugu_tcqh_mmc1.xlsx",
        "sheet": "Downcore T & dD",
        "age_column": "E",
        "temperature_columns": {"author_admitted_maat_c": "F"},
        "admitted_column": "author_admitted_maat_c",
        "latitude": 25.1167,
        "longitude": 98.5667,
        "target": "MAAT",
        "paper_doi": "10.1016/j.scib.2020.11.016",
        "liang_age_column": "CN",
        "liang_temperature_column": "CO",
    },
    "Lugu": {
        "source": SOURCE_ROOT / "zhao2021_lugu_tcqh_mmc1.xlsx",
        "sheet": "Downcore T & dD",
        "age_column": "J",
        "temperature_columns": {"author_admitted_maat_c": "K"},
        "admitted_column": "author_admitted_maat_c",
        "latitude": 27.6833,
        "longitude": 100.75,
        "target": "MAAT",
        "paper_doi": "10.1016/j.scib.2020.11.016",
        "liang_age_column": "DB",
        "liang_temperature_column": "DC",
    },
    "Qinghai": {
        "source": SOURCE_ROOT / "wang2021_qinghai_mmc2.xlsx",
        "sheet": "1. Calculated proxy values",
        "age_column": "A",
        "temperature_columns": {
            "uncorrected_mbt5me_temperature_c": "G",
            "ir6plus7_corrected_temperature_c": "H",
            "author_admitted_irprime6plus7_corrected_maat_c": "I",
        },
        "admitted_column": "author_admitted_irprime6plus7_corrected_maat_c",
        "latitude": 36.8112,
        "longitude": 100.137,
        "target": "MAAT",
        "paper_doi": "10.1016/j.gca.2021.05.004",
        "liang_age_column": "BZ",
        "liang_temperature_column": "CA",
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def finite_number(value: object) -> bool:
    return isinstance(value, (int, float)) and np.isfinite(float(value))


def requested_window(year_ce: float) -> str | None:
    if PRELOCK_START <= year_ce <= PRELOCK_END:
        return "prelock_1500_1920"
    if POSTLOCK_START <= year_ce <= POSTLOCK_END:
        return "postlock_1940_1950"
    return None


def parse_author_record(name: str, definition: dict[str, object]) -> tuple[pd.DataFrame, dict[str, object]]:
    path = Path(definition["source"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[str(definition["sheet"])]
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)
    max_row = int(sheet.max_row or 0)
    rows: list[dict[str, object]] = []
    locked_skipped = 0
    outside_skipped = 0
    decoded_cells = 0
    for row_number in range(2, max_row + 1):
        age_value = sheet[f"{definition['age_column']}{row_number}"].value
        if not finite_number(age_value):
            continue
        age_ka = float(age_value)
        year_ce = 1950.0 - 1000.0 * age_ka
        if LOCKED_START <= year_ce <= LOCKED_END:
            locked_skipped += 1
            continue
        window = requested_window(year_ce)
        if window is None:
            outside_skipped += 1
            continue
        values: dict[str, float | None] = {}
        for field, column in dict(definition["temperature_columns"]).items():
            value = sheet[f"{column}{row_number}"].value
            decoded_cells += 1
            values[str(field)] = float(value) if finite_number(value) else None
        admitted = values[str(definition["admitted_column"])]
        if admitted is None:
            continue
        rows.append(
            {
                "record": name,
                "source_row": row_number,
                "age_ka": age_ka,
                "age_bp": age_ka * 1000.0,
                "year_ce": year_ce,
                "window": window,
                "author_admitted_temperature_c": admitted,
                **values,
            }
        )
    workbook.close()
    frame = pd.DataFrame(rows).sort_values("year_ce").reset_index(drop=True)
    audit = {
        "record": name,
        "source": str(path),
        "source_sha256": sha256(path),
        "locked_rows_skipped_before_temperature_decode": locked_skipped,
        "outside_requested_rows_skipped_before_temperature_decode": outside_skipped,
        "safe_temperature_cells_decoded": decoded_cells,
        "safe_points_retained": int(len(frame)),
    }
    return frame, audit


def parse_liang_record(name: str, definition: dict[str, object]) -> tuple[pd.DataFrame, dict[str, object]]:
    workbook = load_workbook(LIANG_SOURCE, read_only=True, data_only=True)
    sheet = workbook["Data S7.  TP Holocene records"]
    rows: list[dict[str, object]] = []
    locked_skipped = 0
    for row_number in range(3, int(sheet.max_row or 0) + 1):
        age_value = sheet[f"{definition['liang_age_column']}{row_number}"].value
        if not finite_number(age_value):
            continue
        age_bp = float(age_value)
        year_ce = 1950.0 - age_bp
        if LOCKED_START <= year_ce <= LOCKED_END:
            locked_skipped += 1
            continue
        window = requested_window(year_ce)
        if window is None:
            continue
        temperature_value = sheet[f"{definition['liang_temperature_column']}{row_number}"].value
        if not finite_number(temperature_value):
            continue
        rows.append(
            {
                "record": name,
                "source_row": row_number,
                "age_bp": age_bp,
                "year_ce": year_ce,
                "window": window,
                "liang2024_common_rf_maf_c": float(temperature_value),
            }
        )
    workbook.close()
    frame = pd.DataFrame(rows).sort_values("year_ce").reset_index(drop=True)
    return frame, {
        "record": name,
        "locked_rows_skipped_before_temperature_decode": locked_skipped,
        "safe_points_retained": int(len(frame)),
    }


def slope_per_century(years: np.ndarray, values: np.ndarray) -> float | None:
    finite = np.isfinite(years) & np.isfinite(values)
    if finite.sum() < 2 or np.unique(years[finite]).size < 2:
        return None
    return float(np.polyfit(years[finite], values[finite], 1)[0] * 100.0)


def sign(value: float | None) -> str:
    if value is None:
        return "insufficient"
    if value > 1e-12:
        return "warming"
    if value < -1e-12:
        return "cooling"
    return "neutral"


def match_by_year(left: pd.DataFrame, right: pd.DataFrame, tolerance: float = 0.51) -> pd.DataFrame:
    if left.empty or right.empty:
        return pd.DataFrame()
    return pd.merge_asof(
        left.sort_values("year_ce"),
        right.sort_values("year_ce"),
        on="year_ce",
        by="record",
        direction="nearest",
        tolerance=tolerance,
        suffixes=("_author", "_liang"),
    ).dropna(subset=["liang2024_common_rf_maf_c"])


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=False)
    source_paths = sorted({Path(item["source"]) for item in RECORDS.values()} | {LIANG_SOURCE, Path(LINEAGE)})
    hashes_before = {str(path): sha256(path) for path in source_paths}

    author_frames: list[pd.DataFrame] = []
    liang_frames: list[pd.DataFrame] = []
    author_audits: list[dict[str, object]] = []
    liang_audits: list[dict[str, object]] = []
    for name, definition in RECORDS.items():
        author, author_audit = parse_author_record(name, definition)
        liang, liang_audit = parse_liang_record(name, definition)
        author_frames.append(author)
        liang_frames.append(liang)
        author_audits.append(author_audit)
        liang_audits.append(liang_audit)

    author_points = pd.concat(author_frames, ignore_index=True).sort_values(["record", "year_ce"])
    liang_points = pd.concat(liang_frames, ignore_index=True).sort_values(["record", "year_ce"])
    author_points.to_csv(OUTPUT / "original_author_safe_points.csv", index=False)
    liang_points.to_csv(OUTPUT / "liang2024_common_rf_safe_points.csv", index=False)

    lineage = pd.read_csv(LINEAGE)
    lineage = lineage[(lineage.variable == "tas") & lineage.member.isin(PROJECT_MEMBERS)].sort_values("member")
    if set(lineage.member) != set(PROJECT_MEMBERS):
        raise RuntimeError("project-member-lineage-incomplete")

    project: dict[str, dict[str, pd.DataFrame]] = {}
    for name, definition in RECORDS.items():
        project[name] = {
            row.member: read_project_series(Path(row.source_tif), float(definition["latitude"]), float(definition["longitude"]))
            for row in lineage.itertuples(index=False)
        }

    metric_rows: list[dict[str, object]] = []
    match_rows: list[pd.DataFrame] = []
    windows = (
        ("full_safe_prelock", PRELOCK_START, PRELOCK_END),
        ("bridge_1850_1920", 1850.0, PRELOCK_END),
    )
    for name in RECORDS:
        for window_name, start, end in windows:
            author = author_points[(author_points.record == name) & author_points.year_ce.between(start, end)].copy()
            liang = liang_points[(liang_points.record == name) & liang_points.year_ce.between(start, end)].copy()
            author_years = author.year_ce.to_numpy(float)
            author_values = author.author_admitted_temperature_c.to_numpy(float)
            original_trend = slope_per_century(author_years, author_values)
            liang_trend = slope_per_century(
                liang.year_ce.to_numpy(float), liang.liang2024_common_rf_maf_c.to_numpy(float)
            )
            member_trends: list[float] = []
            if len(author_years) >= 2:
                for member in PROJECT_MEMBERS:
                    series = project[name][member]
                    sampled = np.interp(author_years, series.year.to_numpy(float), series.temperature.to_numpy(float))
                    trend = slope_per_century(author_years, sampled)
                    if trend is not None:
                        member_trends.append(trend)
            project_trend = float(np.median(member_trends)) if member_trends else None
            matched = match_by_year(author, liang)
            if not matched.empty:
                matched["window_name"] = window_name
                match_rows.append(matched)
            matched_author_trend = slope_per_century(
                matched.year_ce.to_numpy(float), matched.author_admitted_temperature_c.to_numpy(float)
            ) if not matched.empty else None
            matched_liang_trend = slope_per_century(
                matched.year_ce.to_numpy(float), matched.liang2024_common_rf_maf_c.to_numpy(float)
            ) if not matched.empty else None
            metric_rows.append(
                {
                    "record": name,
                    "window": window_name,
                    "author_target": RECORDS[name]["target"],
                    "author_n": int(len(author)),
                    "author_first_year": float(author_years.min()) if len(author_years) else None,
                    "author_last_year": float(author_years.max()) if len(author_years) else None,
                    "author_trend_c_per_century": original_trend,
                    "project_median_trend_at_author_years_c_per_century": project_trend,
                    "author_project_direction_agreement": sign(original_trend) == sign(project_trend)
                    if original_trend is not None and project_trend is not None
                    else None,
                    "liang_common_rf_n": int(len(liang)),
                    "liang_common_rf_trend_c_per_century": liang_trend,
                    "matched_author_liang_n": int(len(matched)),
                    "matched_author_trend_c_per_century": matched_author_trend,
                    "matched_liang_trend_c_per_century": matched_liang_trend,
                    "matched_calibration_direction_reversal": sign(matched_author_trend) != sign(matched_liang_trend)
                    if matched_author_trend is not None and matched_liang_trend is not None
                    else None,
                }
            )

    metrics = pd.DataFrame(metric_rows)
    metrics.to_csv(OUTPUT / "original_author_window_metrics.csv", index=False)
    matches = pd.concat(match_rows, ignore_index=True) if match_rows else pd.DataFrame()
    matches.to_csv(OUTPUT / "author_liang_matched_points.csv", index=False)

    qinghai = author_points[author_points.record == "Qinghai"].copy()
    qinghai_columns = [
        "uncorrected_mbt5me_temperature_c",
        "ir6plus7_corrected_temperature_c",
        "author_admitted_irprime6plus7_corrected_maat_c",
    ]
    qinghai_sensitivity = []
    for window_name, start, end in windows:
        selected = qinghai[qinghai.year_ce.between(start, end)]
        for column in qinghai_columns:
            qinghai_sensitivity.append(
                {
                    "window": window_name,
                    "calibration_column": column,
                    "n": int(selected[column].notna().sum()),
                    "trend_c_per_century": slope_per_century(
                        selected.year_ce.to_numpy(float), selected[column].to_numpy(float)
                    ),
                }
            )
    pd.DataFrame(qinghai_sensitivity).to_csv(OUTPUT / "qinghai_author_calibration_sensitivity.csv", index=False)

    fig, axes = plt.subplots(1, 3, figsize=(14.8, 4.9), sharey=False)
    colors = {"Tengchongqinghai": "#2563a6", "Lugu": "#27845d", "Qinghai": "#bd4b3f"}
    for ax, name in zip(axes, RECORDS):
        author = author_points[(author_points.record == name) & (author_points.window == "prelock_1500_1920")]
        liang = liang_points[(liang_points.record == name) & (liang_points.window == "prelock_1500_1920")]
        if author.empty:
            continue
        start_year = float(author.year_ce.min())
        author_base = float(author.iloc[0].author_admitted_temperature_c)
        ax.plot(
            author.year_ce,
            author.author_admitted_temperature_c - author_base,
            color="#111111",
            marker="o",
            markersize=4.2,
            linewidth=2.0,
            label="Original author",
            zorder=5,
        )
        if not liang.empty:
            liang_base = float(np.interp(start_year, liang.year_ce, liang.liang2024_common_rf_maf_c))
            ax.plot(
                liang.year_ce,
                liang.liang2024_common_rf_maf_c - liang_base,
                color=colors[name],
                marker="s",
                markersize=3.6,
                linewidth=1.4,
                label="Liang 2024 common RF",
            )
        member_values = []
        years = np.arange(int(np.ceil(start_year)), 1921)
        for member in PROJECT_MEMBERS:
            series = project[name][member]
            values = np.interp(years, series.year, series.temperature)
            base = float(np.interp(start_year, series.year, series.temperature))
            member_values.append(values - base)
        matrix = np.vstack(member_values)
        ax.fill_between(years, np.min(matrix, axis=0), np.max(matrix, axis=0), color="#9ca3af", alpha=0.25)
        ax.plot(years, np.median(matrix, axis=0), color="#555555", linewidth=1.3, label="Project six-member median")
        ax.axhline(0, color="#777777", linewidth=0.7)
        ax.axvspan(1921, 1940, color="#eeeeee", alpha=0.8)
        ax.set_title(name)
        ax.set_xlabel("Year CE")
        ax.grid(color="#e3e3e3", linewidth=0.6)
    axes[0].set_ylabel("Temperature change from first author point (degC)")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=3, frameon=False)
    fig.suptitle("Original author lake reconstructions vs common recalibration and project", y=0.98)
    fig.tight_layout(rect=(0, 0.09, 1, 0.94))
    fig.savefig(OUTPUT / "original_author_common_rf_project_curves.png", dpi=220)
    plt.close(fig)

    for path, source_hash in hashes_before.items():
        if sha256(Path(path)) != source_hash:
            raise RuntimeError(f"source-changed:{path}")
    tif_manifest = [
        {"member": row.member, "path": row.source_tif, "sha256": sha256(Path(row.source_tif))}
        for row in lineage.itertuples(index=False)
    ]
    manifest = {
        "original_sources": [{"path": path, "sha256": value} for path, value in hashes_before.items()],
        "project_members": tif_manifest,
        "driver": {"path": str(DRIVER), "sha256": sha256(DRIVER)},
        "papers": {
            "Zhao2021": "https://doi.org/10.1016/j.scib.2020.11.016",
            "Wang2021": "https://doi.org/10.1016/j.gca.2021.05.004",
            "Liang2024": "https://doi.org/10.1038/s43247-024-01925-3",
        },
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    (OUTPUT / "source_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    metric_records = metrics.astype(object).where(pd.notna(metrics), None).to_dict(orient="records")
    summary = {
        "experiment": "temperature_original_lake_author_reconstructions_v2",
        "status": "complete",
        "method_identity": {
            "Tengchongqinghai": "Zhao et al. 2021 original author Index1-based MAAT",
            "Lugu": "Zhao et al. 2021 original author Index1-based MAAT",
            "Qinghai": "Wang et al. 2021 salinity-corrected MBT'5ME/IR'6+7ME MAAT, workbook column I",
            "Liang2024": "secondary common random-forest MAF recalibration; sensitivity only",
        },
        "metrics": metric_records,
        "author_parse_audits": author_audits,
        "liang_parse_audits": liang_audits,
        "sealed_period_values_used": False,
        "production_modified": False,
        "enkf_run": False,
        "oggm_run": False,
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    (OUTPUT / "summary.json").write_text(json.dumps(summary, indent=2, allow_nan=False) + "\n", encoding="utf-8")

    lines = [
        "# Original author lake-temperature reconstruction audit",
        "",
        "The original author reconstructions are primary site evidence. Liang et al. (2024) is retained only as a common-calibration sensitivity and is not counted as an independent vote.",
        "",
        "| Record | Window | Author n | Author trend | Project trend | Direction agreement | Liang trend | Matched calibration reversal |",
        "|---|---|---:|---:|---:|---|---:|---|",
    ]
    for row in metrics.itertuples(index=False):
        author_trend = "insufficient" if pd.isna(row.author_trend_c_per_century) else f"{row.author_trend_c_per_century:+.3f}"
        project_trend = "insufficient" if pd.isna(row.project_median_trend_at_author_years_c_per_century) else f"{row.project_median_trend_at_author_years_c_per_century:+.3f}"
        liang_trend = "insufficient" if pd.isna(row.liang_common_rf_trend_c_per_century) else f"{row.liang_common_rf_trend_c_per_century:+.3f}"
        agreement = "not evaluated" if pd.isna(row.author_project_direction_agreement) else ("yes" if row.author_project_direction_agreement else "no")
        reversal = "not evaluated" if pd.isna(row.matched_calibration_direction_reversal) else ("yes" if row.matched_calibration_direction_reversal else "no")
        lines.append(
            f"| {row.record} | {row.window} | {row.author_n} | {author_trend} | {project_trend} | {agreement} | {liang_trend} | {reversal} |"
        )
    lines.extend(
        [
            "",
            "The locked interval was excluded from age alone before any temperature cell was decoded. Original and production files were unchanged; EnKF and OGGM were not run; 20CRv3 was not accessed.",
        ]
    )
    (OUTPUT / "FINAL_RESULT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
