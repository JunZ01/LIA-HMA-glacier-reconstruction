from __future__ import annotations

import importlib.util
import math
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "run_round215_haizi_daocheng_sequence_identity_v1.py"
SPEC = importlib.util.spec_from_file_location("round215_v1", SOURCE)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError("Could not load Round 215 v1 driver")
round215 = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(round215)

round215.OUT = round215.LOOP / "round_215_haizi_daocheng_sequence_identity_v3"
round215.CONTRACT = round215.OUT / "ROUND_CONTRACT.md"
round215.DRIVER = Path(__file__).resolve()


def read_haizi_year_first() -> tuple[dict[int, float], int]:
    if round215.sha256(round215.WORKBOOK) != round215.EXPECTED_WORKBOOK_SHA256:
        raise RuntimeError("Haizi workbook hash changed")
    workbook = load_workbook(round215.WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    header = [sheet.cell(1, column).value for column in range(1, 5)]
    if (
        header[0] != "year"
        or header[1] != "std"
        or not isinstance(header[2], str)
        or not header[2].startswith("Temperature(")
        or header[3] != "num"
    ):
        raise RuntimeError("Unexpected Haizi workbook header")
    values: dict[int, float] = {}
    skipped = 0
    for row in sheet.iter_rows(min_row=2, max_col=4, values_only=False):
        year_token = row[0].value
        if year_token is None:
            continue
        year = int(year_token)
        if year < 1115 or year > 2017:
            continue
        if round215.LOCKED_START <= year <= round215.LOCKED_END:
            skipped += 1
            continue
        value = float(row[1].value)
        if not math.isfinite(value):
            raise RuntimeError(f"Non-finite Haizi STD in allowed year {year}")
        values[year] = value
    workbook.close()
    expected = set(range(1115, 1921)) | set(range(1940, 2018))
    if set(values) != expected or skipped != 19:
        raise RuntimeError("Haizi allowed coverage mismatch")
    return values, skipped


round215.read_haizi_year_first = read_haizi_year_first

if __name__ == "__main__":
    round215.main()
