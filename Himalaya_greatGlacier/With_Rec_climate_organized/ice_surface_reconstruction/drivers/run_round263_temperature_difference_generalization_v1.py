from __future__ import annotations

import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt


BAIYU_WORKBOOK_SHA256 = "286e9338e1cfeddb2b51f6805241d1e0989c891af3adbdd81691f1918b794d68"


def project_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / "LIA_PROJECT_FACTS.md").exists():
            return parent
    raise RuntimeError("project-root-not-found")


def fingerprint(path: Path) -> dict[str, object]:
    stat = path.stat()
    return {
        "path": str(path),
        "size_bytes": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def correlation(left: np.ndarray, right: np.ndarray) -> float:
    if len(left) < 2 or np.std(left) == 0 or np.std(right) == 0:
        return float("nan")
    return float(np.corrcoef(left, right)[0, 1])


def fit_ols(x: np.ndarray, y: np.ndarray) -> np.ndarray:
    return np.linalg.lstsq(np.column_stack((np.ones(len(x)), x)), y, rcond=None)[0]


def read_baiyu_modern(workbook_path: Path) -> pd.DataFrame:
    if fingerprint(workbook_path)["sha256"] != BAIYU_WORKBOOK_SHA256:
        raise RuntimeError("baiyu-workbook-hash-changed")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if workbook.sheetnames != ["Tab.1", "Tab.2"]:
        raise RuntimeError("unexpected-baiyu-sheets")
    series: list[dict[int, float]] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        values: dict[int, float] = {}
        for row in sheet.iter_rows(min_row=7, max_col=2, values_only=False):
            year_token = row[0].value
            if year_token is None:
                continue
            year = int(year_token)
            if year < 1962 or year > 2005:
                continue
            value = float(row[1].value)
            if not math.isfinite(value):
                raise RuntimeError(f"nonfinite-baiyu-modern:{sheet_name}:{year}")
            values[year] = value
        if set(values) != set(range(1962, 2006)):
            raise RuntimeError(f"incomplete-baiyu-modern:{sheet_name}")
        series.append(values)
    workbook.close()
    years = np.arange(1962, 2006, dtype=int)
    return pd.DataFrame({
        "year": years,
        "proxy": [series[0][int(year)] for year in years],
        "target": [series[1][int(year)] for year in years],
    })


def evaluate_record(
    record: str,
    record_class: str,
    frame: pd.DataFrame,
    blocks: tuple[tuple[int, int], ...],
) -> tuple[dict[str, object], list[dict[str, object]], pd.DataFrame]:
    frame = frame.sort_values("year").reset_index(drop=True)
    years = frame["year"].to_numpy(int)
    if not np.all(np.diff(years) == 1):
        raise RuntimeError(f"noncontiguous-record:{record}")
    delta_years = years[1:]
    delta_proxy = np.diff(frame["proxy"].to_numpy(float))
    delta_target = np.diff(frame["target"].to_numpy(float))
    oof_prediction = np.full(len(delta_years), np.nan)
    oof_baseline = np.full(len(delta_years), np.nan)
    fold_rows: list[dict[str, object]] = []
    slopes: list[float] = []
    fold_correlations: list[float] = []
    for fold_number, (start, end) in enumerate(blocks, start=1):
        test = (delta_years >= start) & (delta_years <= end)
        train = ~test
        if not test.any() or train.sum() < 2:
            raise RuntimeError(f"invalid-block:{record}:{start}-{end}")
        beta = fit_ols(delta_proxy[train], delta_target[train])
        prediction = beta[0] + beta[1] * delta_proxy[test]
        baseline = np.full(test.sum(), float(delta_target[train].mean()))
        fold_r = correlation(prediction, delta_target[test])
        fold_skill = 1.0 - float(np.sum((delta_target[test] - prediction) ** 2)) / float(
            np.sum((delta_target[test] - baseline) ** 2)
        )
        oof_prediction[test] = prediction
        oof_baseline[test] = baseline
        slopes.append(float(beta[1]))
        fold_correlations.append(fold_r)
        fold_rows.append({
            "record": record,
            "record_class": record_class,
            "fold": fold_number,
            "test_start": start,
            "test_end": end,
            "test_year_count": int(test.sum()),
            "fit_intercept": float(beta[0]),
            "fit_slope": float(beta[1]),
            "test_correlation": fold_r,
            "test_skill_vs_training_mean_delta": fold_skill,
        })
    if not np.isfinite(oof_prediction).all():
        raise RuntimeError(f"incomplete-oof:{record}")
    pooled_sse = float(np.sum((delta_target - oof_prediction) ** 2))
    baseline_sse = float(np.sum((delta_target - oof_baseline) ** 2))
    pooled_r = correlation(oof_prediction, delta_target)
    pooled_skill = 1.0 - pooled_sse / baseline_sse
    positive_fold_fraction = float(np.mean(np.asarray(fold_correlations) > 0))
    gate = bool(
        pooled_r >= 0.30
        and pooled_skill > 0
        and np.all(np.asarray(slopes) > 0)
        and positive_fold_fraction >= 0.75
    )
    metrics = {
        "record": record,
        "record_class": record_class,
        "level_year_start": int(years[0]),
        "level_year_end": int(years[-1]),
        "difference_year_count": int(len(delta_years)),
        "oof_delta_correlation": pooled_r,
        "oof_delta_skill_vs_training_mean": pooled_skill,
        "oof_delta_rmse_degC": float(np.sqrt(pooled_sse / len(delta_years))),
        "positive_slope_fold_count": int(np.sum(np.asarray(slopes) > 0)),
        "fold_count": len(blocks),
        "positive_test_correlation_fold_fraction": positive_fold_fraction,
        "passes_frequency_role_gate": gate,
    }
    curves = pd.DataFrame({
        "record": record,
        "record_class": record_class,
        "difference_year": delta_years,
        "observed_delta_t_degC": delta_target,
        "oof_predicted_delta_t_degC": oof_prediction,
        "oof_baseline_delta_t_degC": oof_baseline,
    })
    return metrics, fold_rows, curves


def main() -> int:
    root = project_root()
    loop = root / "validation_data/results/paleo_surface_validation_v0/climate_reconstruction_improvement_loop_v2"
    output = loop / "round_263_temperature_difference_generalization_v1"
    output.mkdir(parents=True, exist_ok=False)
    paths = {
        "facts": root / "LIA_PROJECT_FACTS.md",
        "driver": Path(__file__).resolve(),
        "round258_curves": loop / "round_258_npl065_precalibration_holdout_v2/precalibration_holdout_curves.csv",
        "round262_summary": loop / "round_262_npl065_first_difference_anchor_psm_v1/summary.json",
        "baiyu_workbook": loop / "round_211_wscp_temperature_data_recovery_v2/official_archive_extracted/Temp_Reconstruction&Meteorological.xlsx",
        "round212_summary": loop / "round_212_wscp_temperature_locked_safe_comparison_v1/summary.json",
        "huang_cru_curves": loop / "round_208_setp_cru_tmin_annual_target_audit_v1/allowed_seasonal_annual_curves.csv",
        "round208_summary": loop / "round_208_setp_cru_tmin_annual_target_audit_v1/summary.json",
        "round214_summary": loop / "round_214_haizi_annual_temperature_data_recovery_v1/summary.json",
        "round226_summary": loop / "round_226_yang2010_setp_622y_metadata_recovery_v2/summary.json",
    }
    before = {name: fingerprint(path) for name, path in paths.items()}

    npl_source = pd.read_csv(paths["round258_curves"])
    npl = npl_source.loc[npl_source["year"].between(1960, 2013), ["year", "signal_free_chronology", "cru_box_djf_tmin"]].rename(
        columns={"signal_free_chronology": "proxy", "cru_box_djf_tmin": "target"}
    )
    baiyu = read_baiyu_modern(paths["baiyu_workbook"])
    huang_source = pd.read_csv(paths["huang_cru_curves"])
    huang = huang_source.loc[
        huang_source["year"].between(1960, 2000),
        ["year", "huang_winter_tmin", "cru_ts401_winter_tmin"],
    ].dropna().rename(columns={"huang_winter_tmin": "proxy", "cru_ts401_winter_tmin": "target"})
    if not np.array_equal(huang["year"].to_numpy(int), np.arange(1960, 2001, dtype=int)):
        raise RuntimeError("incomplete-huang-modern-pair")

    specifications = (
        (
            "NPL065_signal_free",
            "strict_raw_RWL_signal_free_plus_CRU_target",
            npl,
            ((1961, 1969), (1970, 1978), (1979, 1987), (1988, 1996), (1997, 2005), (2006, 2013)),
        ),
        (
            "Baiyu_published_JJ",
            "supporting_published_reconstruction_plus_official_seven_station_target",
            baiyu,
            ((1963, 1973), (1974, 1984), (1985, 1995), (1996, 2005)),
        ),
        (
            "Huang_published_winter_Tmin",
            "supporting_published_reconstruction_plus_exact_CRU_seasonal_target",
            huang,
            ((1961, 1970), (1971, 1980), (1981, 1990), (1991, 2000)),
        ),
    )
    metrics_rows: list[dict[str, object]] = []
    fold_rows: list[dict[str, object]] = []
    curve_frames: list[pd.DataFrame] = []
    for record, record_class, frame, blocks in specifications:
        metrics, folds, curves = evaluate_record(record, record_class, frame, blocks)
        metrics_rows.append(metrics)
        fold_rows.extend(folds)
        curve_frames.append(curves)
    metrics = pd.DataFrame(metrics_rows)
    folds = pd.DataFrame(fold_rows)
    curves = pd.concat(curve_frames, ignore_index=True)
    metrics.to_csv(output / "temperature_difference_generalization_metrics.csv", index=False)
    folds.to_csv(output / "temperature_difference_generalization_fold_metrics.csv", index=False)
    curves.to_csv(output / "temperature_difference_generalization_oof_curves.csv", index=False)

    eligibility = pd.DataFrame([
        {
            "record": "NPL065",
            "raw_or_index_series_available": True,
            "instrumental_target_available": True,
            "signal_free_algorithm_directly_testable": True,
            "role": "strict_method_and_frequency_test",
        },
        {
            "record": "Baiyu",
            "raw_or_index_series_available": False,
            "instrumental_target_available": True,
            "signal_free_algorithm_directly_testable": False,
            "role": "supporting_frequency_test_only_published_reconstruction",
        },
        {
            "record": "Huang",
            "raw_or_index_series_available": False,
            "instrumental_target_available": True,
            "signal_free_algorithm_directly_testable": False,
            "role": "supporting_frequency_test_only_published_reconstruction",
        },
        {
            "record": "Haizi",
            "raw_or_index_series_available": True,
            "instrumental_target_available": False,
            "signal_free_algorithm_directly_testable": False,
            "role": "excluded_Daocheng_numeric_target_not_public_in_supplement",
        },
        {
            "record": "Yang2010",
            "raw_or_index_series_available": False,
            "instrumental_target_available": False,
            "signal_free_algorithm_directly_testable": False,
            "role": "excluded_native_numeric_series_not_public_only_11y_curve_digitized",
        },
    ])
    eligibility.to_csv(output / "temperature_record_method_eligibility.csv", index=False)

    colors = ["#D55E00", "#0072B2", "#009E73"]
    figure, axes = plt.subplots(3, 1, figsize=(12.0, 10.0), constrained_layout=True)
    for axis, (record, _, _, _), color in zip(axes, specifications, colors):
        subset = curves[curves["record"] == record]
        row = metrics.loc[metrics["record"] == record].iloc[0]
        axis.plot(subset["difference_year"], subset["observed_delta_t_degC"], color="#111111", linewidth=1.6, label="Observed delta T")
        axis.plot(subset["difference_year"], subset["oof_predicted_delta_t_degC"], color=color, linewidth=1.35, label=f"Blocked OOF; r={row['oof_delta_correlation']:.3f}, skill={row['oof_delta_skill_vs_training_mean']:.3f}")
        axis.axhline(0, color="#777777", linewidth=0.7)
        axis.set_title(record.replace("_", " "))
        axis.set_ylabel("delta degC")
        axis.grid(axis="y", color="#dddddd", linewidth=0.5)
        axis.spines[["top", "right"]].set_visible(False)
        axis.legend(frameon=False, fontsize=9, ncol=2)
    axes[-1].set_xlabel("Difference year")
    figure.savefig(output / "temperature_difference_generalization_oof.png", dpi=220)
    plt.close(figure)

    after = {name: fingerprint(path) for name, path in paths.items()}
    unchanged = before == after
    if not unchanged:
        raise RuntimeError("source-fingerprint-changed")
    external = metrics[metrics["record"] != "NPL065_signal_free"]
    external_pass_count = int(external["passes_frequency_role_gate"].sum())
    if external_pass_count == 2:
        decision = "frequency_role_supported_by_two_external_records_signal_free_algorithm_remains_npl_specific"
    elif external_pass_count == 1:
        decision = "partial_external_frequency_support_signal_free_algorithm_remains_npl_specific"
    else:
        decision = "no_external_frequency_generalization_retain_npl065_only"
    result = {
        "experiment_id": output.name,
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "status": "complete",
        "frozen_gate": {
            "oof_delta_correlation_minimum": 0.30,
            "oof_delta_skill_minimum_exclusive": 0.0,
            "all_fold_slopes_positive": True,
            "positive_test_correlation_fold_fraction_minimum": 0.75,
        },
        "metrics": metrics_rows,
        "external_pass_count": external_pass_count,
        "decision": decision,
        "interpretation_boundary": "external_records_test_frequency_role_not_signal_free_algorithm_or_absolute_low_frequency_level",
        "locked_1921_1939_values_accessed": False,
        "paleo_proxy_values_accessed": False,
        "jionglacuo_opened_or_hashed": False,
        "rcsigfree_binary_executed": False,
        "enkf_run": False,
        "production_reconstruction_run": False,
        "oggm_run": False,
        "source_fingerprints_unchanged": unchanged,
    }
    (output / "summary.json").write_text(json.dumps(result, indent=2), encoding="ascii")
    (output / "source_manifest.json").write_text(json.dumps({"before": before, "after": after}, indent=2), encoding="ascii")
    lines = [
        "# Round 263 temperature-difference generalization",
        "",
        "Only NPL065 directly tests the raw-RWL signal-free algorithm. Baiyu and Huang test whether the narrower year-to-year frequency role is externally reproducible.",
        "",
    ]
    for row in metrics_rows:
        lines.append(
            f"- {row['record']}: OOF delta r `{row['oof_delta_correlation']:.3f}`, skill `{row['oof_delta_skill_vs_training_mean']:.3f}`, positive-fold fraction `{row['positive_test_correlation_fold_fraction']:.2f}`, pass `{row['passes_frequency_role_gate']}`."
        )
    lines.extend(["", f"Decision: `{decision}`.", ""])
    (output / "FINAL_RESULT.md").write_text("\n".join(lines), encoding="ascii")
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
