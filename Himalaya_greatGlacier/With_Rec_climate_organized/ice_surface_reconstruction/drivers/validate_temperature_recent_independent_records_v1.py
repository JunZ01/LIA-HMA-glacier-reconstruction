"""Blind validation of the frozen HMA increment against recent eastern records."""

from __future__ import annotations

import json
from pathlib import Path

import h5py
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

import validate_temperature_external_low_frequency_v1 as base
import validate_temperature_external_low_frequency_v2 as coverage


SOURCE = base.ROOT / "validation_data/sources/external_hma_temperature_candidates_round2_v1"
OUTPUT = (
    base.RESULTS
    / "climate_reconstruction_temperature_recent_independent_blind_validation_v1"
)
PUBLIC = Path(
    r"E:\3spyderProject\HMA_500\Down_scale_new\out_data\climate_anomaly_05deg_with_ensemble_1000-2000.nc"
)
VARIABLE = "tas_ensemble_mean_rec_result"
YEARS = np.arange(1000, 2001)


def read_workbook_series(
    path: Path,
    value_column: int,
    first_data_row: int,
    valid_start: int,
    valid_end: int,
) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = []
    for row in sheet.iter_rows(min_row=first_data_row):
        if not row:
            continue
        try:
            year = int(float(row[0].value))
        except (TypeError, ValueError):
            continue
        if base.LOCKED_START <= year <= base.LOCKED_END:
            continue
        if not valid_start <= year <= valid_end:
            continue
        value = row[value_column - 1].value
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        rows.append({"year": year, "external": float(value)})
    workbook.close()
    return pd.DataFrame(rows).sort_values("year").reset_index(drop=True)


def field_without_locked(dataset: h5py.Dataset) -> np.ndarray:
    result = np.full(dataset.shape, np.nan, dtype=float)
    result[:921] = np.asarray(dataset[:921], dtype=float)
    result[940:] = np.asarray(dataset[940:], dtype=float)
    return result


def box_series(
    field: np.ndarray,
    lat: np.ndarray,
    lon: np.ndarray,
    bounds: tuple[float, float, float, float],
) -> tuple[np.ndarray, int]:
    lon_min, lon_max, lat_min, lat_max = bounds
    mask = (
        (lat[:, None] >= lat_min)
        & (lat[:, None] <= lat_max)
        & (lon[None, :] >= lon_min)
        & (lon[None, :] <= lon_max)
    )
    weights = np.cos(np.deg2rad(lat))[:, None] * mask
    values = np.where(mask[None, :, :], field, np.nan)
    valid = np.isfinite(values)
    numerator = np.nansum(values * weights[None, :, :], axis=(1, 2))
    denominator = np.sum(valid * weights[None, :, :], axis=(1, 2))
    series = np.divide(
        numerator, denominator, out=np.full(len(field), np.nan), where=denominator > 0
    )
    return series, int(mask.sum())


def evaluate(
    local: pd.DataFrame,
    external: pd.DataFrame,
    label: str,
    modern_start: int,
    modern_end: int,
) -> tuple[dict[str, object], pd.DataFrame]:
    merged = local.merge(external, on="year", how="inner")

    def zscore(column: str) -> pd.Series:
        modern = merged.loc[merged.year.between(modern_start, modern_end), column].dropna()
        if len(modern) < 25 or modern.std(ddof=0) <= 0:
            raise ValueError(f"Insufficient modern values for {label}:{column}")
        return (merged[column] - modern.mean()) / modern.std(ddof=0)

    for column in ["project", "candidate", "external"]:
        merged[f"{column}_z"] = zscore(column)

    rows = []
    for start, end in base.WINDOWS:
        subset = merged.loc[merged.year.between(start, end)]
        counts = subset[["project_z", "candidate_z", "external_z"]].notna().sum()
        if counts.min() < 35:
            continue
        rows.append(
            {
                "record": label,
                "window_start": start,
                "window_end": end,
                "project_z": subset.project_z.mean(),
                "candidate_z": subset.candidate_z.mean(),
                "external_z": subset.external_z.mean(),
                "n_external": int(counts.external_z),
            }
        )
    windows = pd.DataFrame(rows)
    if len(windows) != 8:
        raise ValueError(f"{label} retained {len(windows)} rather than 8 windows")
    project_rmse = float(np.sqrt(np.mean((windows.project_z - windows.external_z) ** 2)))
    candidate_rmse = float(
        np.sqrt(np.mean((windows.candidate_z - windows.external_z) ** 2))
    )
    improvement = (project_rmse - candidate_rmse) / project_rmse
    return (
        {
            "record": label,
            "modern_start": modern_start,
            "modern_end": modern_end,
            "modern_n": int(
                merged.loc[merged.year.between(modern_start, modern_end), "external"].notna().sum()
            ),
            "project_50year_rmse_z": project_rmse,
            "candidate_50year_rmse_z": candidate_rmse,
            "rmse_improvement_fraction": improvement,
            "window_correlation_project": float(windows[["project_z", "external_z"]].corr().iloc[0, 1]),
            "window_correlation_candidate": float(windows[["candidate_z", "external_z"]].corr().iloc[0, 1]),
        },
        windows,
    )


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contract = OUTPUT / "ROUND_CONTRACT.md"
    if not contract.exists():
        raise FileNotFoundError("Contract must exist before numeric parsing")
    paths = {
        "jionglacuo": SOURCE / "Temperature reconstruction.xlsx",
        "haizi": SOURCE / "41598_2024_79096_MOESM1_ESM.xlsx",
        "public_field": PUBLIC,
        "frozen_candidate": base.CANDIDATE,
        "download_manifest": SOURCE / "source_manifest.json",
        "contract": contract,
    }
    before = {name: base.sha256(path) for name, path in paths.items()}

    jion = read_workbook_series(paths["jionglacuo"], 2, 1, 1210, 2023)
    haizi = read_workbook_series(paths["haizi"], 3, 2, 1115, 2017)

    with h5py.File(PUBLIC, "r") as handle:
        years = np.asarray(handle["time"], dtype=int).ravel()
        lat = np.asarray(handle["lat"], dtype=float).ravel()
        lon = np.asarray(handle["lon"], dtype=float).ravel()
        if not np.array_equal(years, YEARS):
            raise ValueError("Unexpected public-field years")
        field = field_without_locked(handle[VARIABLE])

    boxes = {
        "Haizi_annual": (99.75, 101.25, 28.75, 30.25),
        "Jionglacuo_JJAS": (93.5, 95.5, 30.0, 32.0),
    }
    local = pd.DataFrame({"year": YEARS})
    cell_counts = {}
    for label, bounds in boxes.items():
        local[label], cell_counts[label] = box_series(field, lat, lon, bounds)

    frozen = base.read_candidate(base.CANDIDATE)
    frozen["increment"] = frozen.candidate - frozen.project
    local = local.merge(frozen[["year", "increment"]], on="year", how="left")

    definitions = [
        ("Haizi_annual", haizi, 1959, 2000),
        ("Jionglacuo_JJAS", jion, 1970, 2000),
    ]
    metrics = []
    window_tables = []
    for label, external, modern_start, modern_end in definitions:
        control = local[["year", label, "increment"]].rename(columns={label: "project"})
        control["candidate"] = control.project + control.increment
        metric, windows = evaluate(control, external, label, modern_start, modern_end)
        metric["grid_cell_count"] = cell_counts[label]
        metrics.append(metric)
        window_tables.append(windows)

    metrics_table = pd.DataFrame(metrics)
    windows_table = pd.concat(window_tables, ignore_index=True)
    haizi_improvement = float(
        metrics_table.loc[metrics_table.record == "Haizi_annual", "rmse_improvement_fraction"].iloc[0]
    )
    jion_improvement = float(
        metrics_table.loc[metrics_table.record == "Jionglacuo_JJAS", "rmse_improvement_fraction"].iloc[0]
    )
    gates = {
        "haizi_annual_rmse_improves_at_least_10pct": bool(haizi_improvement >= 0.10 - 1e-12),
        "jionglacuo_summer_does_not_worsen_more_than_5pct": bool(jion_improvement >= -0.05 - 1e-12),
    }
    passes = bool(all(gates.values()))

    metrics_table.to_csv(OUTPUT / "recent_record_metrics.csv", index=False)
    windows_table.to_csv(OUTPUT / "recent_record_50year_windows.csv", index=False)

    fig, axes = plt.subplots(1, 2, figsize=(13, 5), constrained_layout=True)
    for axis, (label, _, _, _) in zip(axes, definitions):
        table = windows_table.loc[windows_table.record == label]
        x = (table.window_start + table.window_end) / 2
        axis.plot(x, table.external_z, color="#111111", marker="o", label="External")
        axis.plot(x, table.project_z, color="#d55e00", marker="s", label="Project")
        axis.plot(x, table.candidate_z, color="#0072b2", marker="^", label="Frozen increment")
        axis.axhline(0, color="#888888", linewidth=0.8)
        axis.set_title(label)
        axis.set_xlabel("50-year window midpoint")
        axis.set_ylabel("Modern-standardized temperature")
        axis.grid(alpha=0.2)
    axes[0].legend()
    fig.savefig(OUTPUT / "recent_independent_temperature_curves.png", dpi=220)
    plt.close(fig)

    after = {name: base.sha256(path) for name, path in paths.items()}
    (OUTPUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "inputs": {
                    name: {"path": str(paths[name]), "sha256_before": before[name], "sha256_after": after[name]}
                    for name in paths
                },
                "all_input_hashes_unchanged": before == after,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    result = {
        "status": "complete",
        "haizi_annual_rmse_improvement_fraction": haizi_improvement,
        "jionglacuo_jjas_rmse_improvement_fraction": jion_improvement,
        "gates": gates,
        "passes_recent_eastern_hma_record_gates": passes,
        "locked_1921_1939_workbook_rows_skipped_before_value_access": True,
        "locked_1921_1939_field_slices_read": False,
        "phase_1901_1920_used": False,
        "western_or_central_hma_validated": False,
        "spatial_field_release_authorized": False,
        "production_release_authorized": False,
        "decision": (
            "recent_eastern_records_support_frozen_increment_but_do_not_release_field"
            if passes
            else "recent_eastern_records_do_not_jointly_support_frozen_increment"
        ),
    }
    (OUTPUT / "summary.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    (OUTPUT / "FINAL_RESULT.md").write_text(
        "# Recent Independent Temperature Blind Validation\n\n"
        f"- Haizi annual RMSE improvement: `{100*haizi_improvement:.2f}%`.\n"
        f"- Jionglacuo JJAS RMSE improvement: `{100*jion_improvement:.2f}%`.\n"
        f"- Both frozen gates pass: `{str(passes).lower()}`.\n"
        "- Eastern-HMA evidence only; no spatial or production field is released.\n",
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
