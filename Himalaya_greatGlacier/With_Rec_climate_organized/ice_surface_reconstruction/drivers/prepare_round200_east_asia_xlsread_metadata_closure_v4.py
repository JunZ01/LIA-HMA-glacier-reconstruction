from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
SOURCE = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_199_east_asia_proxy_identity_metadata_v2/Proxies_Ensemble.xlsx"
)
OUT = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_200_east_asia_xlsread_metadata_closure_v4"
)


def main() -> None:
    if OUT.exists():
        raise FileExistsError(f"Refusing to overwrite independent result: {OUT}")
    OUT.mkdir(parents=True)

    source = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = source[source.sheetnames[0]]
    top = list(sheet.iter_rows(min_row=1, max_row=14, min_col=1, max_col=424, values_only=True))
    source.close()
    if len(top) != 14 or any(len(row) != 424 for row in top):
        raise RuntimeError("Unexpected source metadata dimensions")

    synthetic = Workbook(write_only=False)
    target = synthetic.active
    target.title = "Sheet1"
    for row_index, row in enumerate(top, start=1):
        for column_index, value in enumerate(row, start=1):
            if value is not None:
                target.cell(row=row_index, column=column_index, value=value)
    for offset, value in enumerate(range(1, 2020), start=15):
        target.cell(row=offset, column=1, value=value)
    path = OUT / "metadata_plus_year_labels_only.xlsx"
    synthetic.save(path)

    manifest = {
        "source_path": str(SOURCE),
        "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "source_rows_decoded": "1-14 only",
        "synthetic_path": str(path),
        "synthetic_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "synthetic_annual_content": "column A year labels 1-2019; all proxy cells blank",
        "original_annual_proxy_values_decoded": False,
        "locked_period_values_accessed": False,
        "jionglacuo_accessed": False,
    }
    (OUT / "preparation_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()

