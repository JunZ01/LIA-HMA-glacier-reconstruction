"""Blocked out-of-fold annual-temperature PSM test for Chongce delta-18O."""

from __future__ import annotations

import json
from pathlib import Path

import h5py
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

import validate_temperature_external_low_frequency_v1 as base


SOURCE = (
    base.ROOT
    / "validation_data/sources/external_hma_temperature_candidates_lake_ice_v1"
    / "chongce_zenodo_5766235.xlsx"
)
PUBLIC = Path(
    r"E:\3spyderProject\HMA_500\Down_scale_new\out_data\climate_anomaly_05deg_with_ensemble_1000-2000.nc"
)
OUTPUT = base.RESULTS / "climate_reconstruction_temperature_chongce_psm_oof_v1"
BLOCKS = [(1940, 1951), (1952, 1963), (1964, 1975), (1976, 1987), (1988, 2000)]
PRODUCTS = {
    "Berkeley": "tas_obs_bea_05deg",
    "CRU": "tas_obs_cru_tmp_05deg",
    "GISTEMP": "tas_obs_GIST_05deg",
    "NOAA": "tas_obs_NOAA_05deg",
    "HadCRUT5": "tas_obs_had5_05deg",
}


def read_modern_chongce(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = []
    for row in sheet.iter_rows(min_row=6):
        try:
            year = int(float(row[0].value))
        except (TypeError, ValueError):
            continue
        if not 1940 <= year <= 2000:
            continue
        isotope = row[1].value
        reconstructed = row[2].value
        if not isinstance(isotope, (int, float)) or not isinstance(reconstructed, (int, float)):
            continue
        rows.append(
            {"year": year, "delta18O": float(isotope), "published_temperature": float(reconstructed)}
        )
    workbook.close()
    result = pd.DataFrame(rows).sort_values("year").reset_index(drop=True)
    if not np.array_equal(result.year.to_numpy(), np.arange(1940, 2001)):
        raise ValueError("Chongce modern coverage is not complete 1940-2000")
    return result


def area_series(field: np.ndarray, lat: np.ndarray, lon: np.ndarray) -> np.ndarray:
    mask = (
        (lat[:, None] >= 34.75)
        & (lat[:, None] <= 35.75)
        & (lon[None, :] >= 80.75)
        & (lon[None, :] <= 81.75)
    )
    weights = np.cos(np.deg2rad(lat))[:, None] * mask
    values = np.where(mask[None, :, :], field, np.nan)
    valid = np.isfinite(values)
    numerator = np.nansum(values * weights[None, :, :], axis=(1, 2))
    denominator = np.sum(valid * weights[None, :, :], axis=(1, 2))
    return np.divide(
        numerator, denominator, out=np.full(field.shape[0], np.nan), where=denominator > 0
    )


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contract = OUTPUT / "ROUND_CONTRACT.md"
    paths = {"chongce": SOURCE, "public_field": PUBLIC, "contract": contract}
    before = {name: base.sha256(path) for name, path in paths.items()}
    proxy = read_modern_chongce(SOURCE)

    with h5py.File(PUBLIC, "r") as handle:
        all_years = np.asarray(handle["time"], dtype=int).ravel()
        selected = np.flatnonzero((all_years >= 1940) & (all_years <= 2000))
        if not np.array_equal(all_years[selected], np.arange(1940, 2001)):
            raise ValueError("Modern public-field indexing failed")
        lat = np.asarray(handle["lat"], dtype=float).ravel()
        lon = np.asarray(handle["lon"], dtype=float).ravel()
        observations = {
            name: area_series(np.asarray(handle[variable][selected], dtype=float), lat, lon)
            for name, variable in PRODUCTS.items()
        }

    years = proxy.year.to_numpy(int)
    predictor = proxy.delta18O.to_numpy(float)
    rows = []
    predictions = []
    for product, target in observations.items():
        valid_full = np.isfinite(target) & np.isfinite(predictor)
        full_coef = np.linalg.lstsq(
            np.column_stack([np.ones(valid_full.sum()), predictor[valid_full]]),
            target[valid_full],
            rcond=None,
        )[0]
        oof = np.full(len(years), np.nan)
        baseline = np.full(len(years), np.nan)
        fold_improvements = []
        for fold, (start, end) in enumerate(BLOCKS, start=1):
            test = (years >= start) & (years <= end)
            train = ~test
            valid_train = train & np.isfinite(target) & np.isfinite(predictor)
            valid_test = test & np.isfinite(target) & np.isfinite(predictor)
            if valid_train.sum() < 35 or valid_test.sum() < 8:
                raise ValueError(f"Insufficient {product} support in fold {fold}")
            coef = np.linalg.lstsq(
                np.column_stack([np.ones(valid_train.sum()), predictor[valid_train]]),
                target[valid_train],
                rcond=None,
            )[0]
            oof[valid_test] = coef[0] + coef[1] * predictor[valid_test]
            baseline[valid_test] = np.mean(target[valid_train])
            model_rmse = float(np.sqrt(np.mean((target[valid_test] - oof[valid_test]) ** 2)))
            intercept_rmse = float(
                np.sqrt(np.mean((target[valid_test] - baseline[valid_test]) ** 2))
            )
            improvement = 1.0 - model_rmse / intercept_rmse
            fold_improvements.append(improvement)
            rows.append(
                {
                    "product": product,
                    "fold": fold,
                    "test_start": start,
                    "test_end": end,
                    "train_n": int(valid_train.sum()),
                    "test_n": int(valid_test.sum()),
                    "coefficient_delta18O": float(coef[1]),
                    "intercept_rmse_degC": intercept_rmse,
                    "model_rmse_degC": model_rmse,
                    "rmse_improvement_fraction": improvement,
                }
            )
        valid_oof = np.isfinite(oof) & np.isfinite(target)
        pooled_corr = float(np.corrcoef(oof[valid_oof], target[valid_oof])[0, 1])
        predictions.append(
            {
                "product": product,
                "full_coefficient_delta18O": float(full_coef[1]),
                "median_fold_rmse_improvement": float(np.median(fold_improvements)),
                "pooled_oof_correlation": pooled_corr,
                "oof_prediction": oof,
                "target": target,
            }
        )

    folds = pd.DataFrame(rows)
    summary_rows = [
        {key: value for key, value in item.items() if key not in {"oof_prediction", "target"}}
        for item in predictions
    ]
    products = pd.DataFrame(summary_rows)
    folds.to_csv(OUTPUT / "chongce_psm_fold_metrics.csv", index=False)
    products.to_csv(OUTPUT / "chongce_psm_product_summary.csv", index=False)

    improvements = products.median_fold_rmse_improvement.to_numpy(float)
    correlations = products.pooled_oof_correlation.to_numpy(float)
    coefficients = products.full_coefficient_delta18O.to_numpy(float)
    gates = {
        "all_five_full_coefficients_positive": bool(np.all(coefficients > 0)),
        "at_least_four_products_positive_median_improvement": bool(np.sum(improvements > 0) >= 4),
        "median_product_improvement_at_least_5pct": bool(np.median(improvements) >= 0.05),
        "median_pooled_oof_correlation_at_least_0p20": bool(np.median(correlations) >= 0.20),
    }
    passes = bool(all(gates.values()))

    fig, axes = plt.subplots(2, 1, figsize=(12, 8), constrained_layout=True)
    for item in predictions:
        axes[0].plot(years, item["target"], linewidth=1.0, alpha=0.55, label=item["product"])
        axes[1].plot(years, item["oof_prediction"], linewidth=1.2, label=item["product"])
    axes[0].set(title="Local observed annual temperature products", ylabel="Temperature anomaly (degC)")
    axes[1].set(title="Chongce delta-18O blocked OOF predictions", xlabel="Year", ylabel="Predicted anomaly (degC)")
    for axis in axes:
        axis.grid(alpha=0.2)
        axis.legend(ncol=3, fontsize=8)
    fig.savefig(OUTPUT / "chongce_psm_oof.png", dpi=220)
    plt.close(fig)

    after = {name: base.sha256(path) for name, path in paths.items()}
    (OUTPUT / "source_manifest.json").write_text(
        json.dumps({
            "inputs": {
                name: {"path": str(paths[name]), "sha256_before": before[name], "sha256_after": after[name]}
                for name in paths
            },
            "all_input_hashes_unchanged": before == after,
        }, indent=2), encoding="utf-8"
    )
    result = {
        "status": "complete",
        "median_product_rmse_improvement_fraction": float(np.median(improvements)),
        "positive_improvement_product_count": int(np.sum(improvements > 0)),
        "median_pooled_oof_correlation": float(np.median(correlations)),
        "full_coefficient_positive_product_count": int(np.sum(coefficients > 0)),
        "gates": gates,
        "passes_chongce_psm_gates": passes,
        "read_period": "1940-2000_only",
        "phase_1901_1920_used": False,
        "locked_1921_1939_accessed": False,
        "enkf_pilot_authorized": passes,
        "current_reconstruction_modified": False,
    }
    (OUTPUT / "summary.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    (OUTPUT / "FINAL_RESULT.md").write_text(
        "# Chongce Annual-Temperature PSM OOF\n\n"
        f"- Median product RMSE improvement: `{100*np.median(improvements):.2f}%`.\n"
        f"- Positive-improvement products: `{int(np.sum(improvements > 0))}/5`.\n"
        f"- Median pooled OOF correlation: `{np.median(correlations):.3f}`.\n"
        f"- All gates pass: `{str(passes).lower()}`.\n"
        f"- Separate EnKF pilot authorized: `{str(passes).lower()}`.\n",
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
