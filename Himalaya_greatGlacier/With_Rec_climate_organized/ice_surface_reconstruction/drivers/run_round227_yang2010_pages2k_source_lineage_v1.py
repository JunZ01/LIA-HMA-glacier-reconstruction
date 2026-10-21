from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
LOOP = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2"
)
ROUND226 = LOOP / "round_226_yang2010_setp_622y_metadata_recovery_v2"
OUT = LOOP / "round_227_yang2010_pages2k_source_lineage_v1"
FACTS = ROOT / "LIA_PROJECT_FACTS.md"
DRIVER = Path(__file__).resolve()
CONTRACT = DRIVER.with_name("round227_yang2010_pages2k_source_lineage_v1_contract.md")
SUPPLEMENT_URL = (
    "https://www.ebi.ac.uk/europepmc/webservices/rest/"
    "PMC5505119/supplementaryFiles"
)
TARGET_TERMS = (
    "asia_222",
    "asia_223",
    "mainjt",
    "nyinjt",
    "10.1177/0959683609350388",
    "622-year regional temperature history",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def download(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "LIA-HMA-metadata-audit/1.0",
            "Accept": "application/zip",
        },
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        payload = response.read()
    if not payload.startswith(b"PK"):
        raise RuntimeError("Europe PMC response is not a ZIP archive")
    return payload


def normalized(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def unique_headers(values: tuple[object, ...]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = normalized(value) or f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        suffix = f"__{counts[base]}" if counts[base] > 1 else ""
        headers.append(base + suffix)
    return headers


def extract_matches(workbook_path: Path) -> tuple[list[dict[str, object]], list[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    matches: list[dict[str, object]] = []
    sheet_names = list(workbook.sheetnames)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                first = next(rows)
            except StopIteration:
                continue
            headers = unique_headers(first)
            for excel_row, values in enumerate(rows, start=2):
                text = " | ".join(normalized(value) for value in values).lower()
                terms = [term for term in TARGET_TERMS if term in text]
                if not terms:
                    continue
                row = {
                    header: normalized(values[index]) if index < len(values) else ""
                    for index, header in enumerate(headers)
                }
                matches.append(
                    {
                        "sheet": sheet.title,
                        "excel_row": excel_row,
                        "matched_terms": terms,
                        "metadata": row,
                    }
                )
    finally:
        workbook.close()
    return matches, sheet_names


def flatten_matches(matches: list[dict[str, object]]) -> list[dict[str, str]]:
    flattened: list[dict[str, str]] = []
    all_fields: list[str] = []
    for match in matches:
        for field in match["metadata"]:
            if field not in all_fields:
                all_fields.append(field)
    for match in matches:
        metadata = match["metadata"]
        flattened.append(
            {
                "sheet": str(match["sheet"]),
                "excel_row": str(match["excel_row"]),
                "matched_terms": ";".join(match["matched_terms"]),
                **{field: str(metadata.get(field, "")) for field in all_fields},
            }
        )
    return flattened


def classify(matches: list[dict[str, object]]) -> dict[str, object]:
    target_rows = {
        term: [match for match in matches if term in match["matched_terms"]]
        for term in TARGET_TERMS[:4]
    }
    yang_rows = [
        match
        for match in matches
        if any(term in match["matched_terms"] for term in TARGET_TERMS[4:])
    ]
    all_target_text = " ".join(
        " ".join(str(value).lower() for value in match["metadata"].values())
        for term in TARGET_TERMS[:4]
        for match in target_rows[term]
    )
    yang_doi_in_target_rows = TARGET_TERMS[4] in all_target_text
    if not any(target_rows.values()):
        decision = "target_identifiers_not_found_in_official_supplement"
    elif yang_doi_in_target_rows:
        decision = "yang2010_publication_overlap_confirmed"
    else:
        decision = "yang2010_not_named_in_target_metadata_sample_overlap_requires_source_review"
    return {
        "target_match_counts": {term: len(rows) for term, rows in target_rows.items()},
        "yang_match_count": len(yang_rows),
        "yang_doi_in_target_rows": yang_doi_in_target_rows,
        "decision": decision,
    }


def main() -> None:
    if OUT.exists():
        raise FileExistsError(f"Refusing to overwrite independent result: {OUT}")
    required = [FACTS, DRIVER, CONTRACT, ROUND226 / "summary.json"]
    if missing := [str(path) for path in required if not path.exists()]:
        raise FileNotFoundError("Missing required input: " + "; ".join(missing))
    facts_before = sha256(FACTS)
    OUT.mkdir(parents=True)

    payload = download(SUPPLEMENT_URL)
    archive_path = OUT / "PMC5505119_SupplementaryFiles.zip"
    archive_path.write_bytes(payload)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        candidates = [
            name for name in names if name.lower().endswith("sdata201788-s3.xlsx")
        ]
        if len(candidates) != 1:
            raise RuntimeError(f"Expected one metadata workbook, found: {candidates}")
        workbook_member = candidates[0]
        workbook_bytes = archive.read(workbook_member)
    if not workbook_bytes.startswith(b"PK"):
        raise RuntimeError("Supplementary metadata workbook is not XLSX")
    workbook_path = OUT / "sdata201788-s3.xlsx"
    workbook_path.write_bytes(workbook_bytes)

    matches, sheet_names = extract_matches(workbook_path)
    classification = classify(matches)
    (OUT / "matched_metadata_rows.json").write_text(
        json.dumps(matches, indent=2, ensure_ascii=True) + "\n", encoding="utf-8"
    )
    flat = flatten_matches(matches)
    if flat:
        with (OUT / "matched_metadata_rows.csv").open(
            "w", encoding="utf-8", newline=""
        ) as stream:
            writer = csv.DictWriter(stream, fieldnames=list(flat[0]))
            writer.writeheader()
            writer.writerows(flat)

    facts_after = sha256(FACTS)
    if facts_after != facts_before:
        raise RuntimeError("Canonical facts changed during read-only audit")
    summary = {
        "round": 227,
        "version": 1,
        "source": {
            "url": SUPPLEMENT_URL,
            "archive_sha256": sha256(archive_path),
            "archive_size_bytes": archive_path.stat().st_size,
            "archive_members": names,
            "metadata_member": workbook_member,
            "metadata_workbook_sha256": sha256(workbook_path),
            "metadata_workbook_size_bytes": workbook_path.stat().st_size,
            "sheet_names": sheet_names,
        },
        "classification": classification,
        "matched_row_count": len(matches),
        "annual_values_read": False,
        "locked_period_values_read": False,
        "jionglacuo_accessed": False,
        "production_files_modified": False,
        "facts_hash_unchanged": facts_after == facts_before,
    }
    (OUT / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=True) + "\n", encoding="utf-8"
    )
    shutil.copy2(CONTRACT, OUT / "ROUND_CONTRACT.md")
    (OUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "facts": {"path": str(FACTS), "sha256": facts_after},
                "round226_summary": {
                    "path": str(ROUND226 / "summary.json"),
                    "sha256": sha256(ROUND226 / "summary.json"),
                },
                "driver": {"path": str(DRIVER), "sha256": sha256(DRIVER)},
                "contract": {"path": str(CONTRACT), "sha256": sha256(CONTRACT)},
                "official_zip": {"path": str(archive_path), "sha256": sha256(archive_path)},
                "official_metadata_workbook": {
                    "path": str(workbook_path),
                    "sha256": sha256(workbook_path),
                },
            },
            indent=2,
            ensure_ascii=True,
        )
        + "\n",
        encoding="utf-8",
    )
    (OUT / "FINAL_RESULT.md").write_text(
        "# Round 227 PAGES2k source-lineage result\n\n"
        f"- Official metadata matches: `{len(matches)}` rows.\n"
        f"- Target counts: `{classification['target_match_counts']}`.\n"
        f"- Yang DOI appears in target rows: `{classification['yang_doi_in_target_rows']}`.\n"
        f"- Decision: `{classification['decision']}`.\n"
        "- No annual values, locked-period values, or Jionglacuo material were read.\n",
        encoding="utf-8",
    )
    print(json.dumps(summary, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
