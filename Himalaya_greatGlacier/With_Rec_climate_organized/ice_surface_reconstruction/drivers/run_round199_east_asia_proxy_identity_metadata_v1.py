from __future__ import annotations

import csv
import hashlib
import json
import re
import struct
import urllib.request
import zlib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
OUT = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_199_east_asia_proxy_identity_metadata_v1"
)
ROUND197_METADATA = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_197_east_asia_temperature_metadata_v2/zenodo_record_18477496.json"
)
ROUND198_INVENTORY = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_198_east_asia_remote_data_inventory_v1/data_zip_central_directory.csv"
)
PROJECT_LEDGER = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_108_proxy_metadata_reconciliation_v2/proxy_evidence_master_ledger_reconciled.csv"
)
MEMBER = "Data (2)/Proxies_Ensemble.xlsx"
EXPECTED_ARCHIVE_SIZE = 1_177_089_719


def get_range(url: str, start: int, end: int) -> bytes:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "LIA-HMA-targeted-member-audit/1.0",
            "Range": f"bytes={start}-{end}",
            "Accept-Encoding": "identity",
        },
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        if getattr(response, "status", None) != 206:
            raise RuntimeError(f"Range request not honored: {getattr(response, 'status', None)}")
        payload = response.read()
    if len(payload) != end - start + 1:
        raise RuntimeError("Range response length mismatch")
    return payload


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def as_float(value: object) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def as_int(value: object) -> int | None:
    number = as_float(value)
    return int(round(number)) if number is not None else None


def main() -> None:
    if OUT.exists():
        raise FileExistsError(f"Refusing to overwrite independent result: {OUT}")
    OUT.mkdir(parents=True)

    metadata_bytes = ROUND197_METADATA.read_bytes()
    metadata = json.loads(metadata_bytes.decode("utf-8"))
    data_item = next(item for item in metadata["files"] if item["key"] == "Data.zip")
    if int(data_item["size"]) != EXPECTED_ARCHIVE_SIZE:
        raise RuntimeError("Data.zip size changed")
    url = data_item["links"].get("content", data_item["links"].get("self"))

    with ROUND198_INVENTORY.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    selected = [row for row in rows if row["member"] == MEMBER]
    if len(selected) != 1:
        raise RuntimeError(f"Expected one proxy workbook member, found {len(selected)}")
    entry = selected[0]
    local_offset = int(entry["local_header_offset"])
    compressed_size = int(entry["compressed_size"])
    uncompressed_size = int(entry["uncompressed_size"])
    expected_crc = int(entry["crc32"], 16)
    method = int(entry["compression_method"])

    local = get_range(url, local_offset, local_offset + 29)
    values = struct.unpack("<4s5H3I2H", local)
    if values[0] != b"PK\x03\x04":
        raise RuntimeError("Invalid local file header signature")
    name_len, extra_len = values[-2:]
    header_tail = get_range(url, local_offset + 30, local_offset + 30 + name_len + extra_len - 1)
    member_name = header_tail[:name_len].decode("utf-8" if values[2] & 0x800 else "cp437")
    if member_name != MEMBER:
        raise RuntimeError(f"Local member mismatch: {member_name}")
    data_start = local_offset + 30 + name_len + extra_len
    compressed = get_range(url, data_start, data_start + compressed_size - 1)
    if method == 8:
        workbook_bytes = zlib.decompress(compressed, -15)
    elif method == 0:
        workbook_bytes = compressed
    else:
        raise RuntimeError(f"Unsupported ZIP compression method: {method}")
    if len(workbook_bytes) != uncompressed_size:
        raise RuntimeError("Uncompressed workbook size mismatch")
    if zlib.crc32(workbook_bytes) & 0xFFFFFFFF != expected_crc:
        raise RuntimeError("Workbook CRC mismatch")

    workbook_path = OUT / "Proxies_Ensemble.xlsx"
    workbook_path.write_bytes(workbook_bytes)
    workbook_sha256 = hashlib.sha256(workbook_bytes).hexdigest()

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise RuntimeError(f"Expected one sheet, found {workbook.sheetnames}")
    sheet = workbook[workbook.sheetnames[0]]
    metadata_rows = list(sheet.iter_rows(min_row=1, max_row=14, values_only=True))
    workbook.close()
    if len(metadata_rows) != 14 or len(metadata_rows[0]) < 2:
        raise RuntimeError("Unexpected proxy workbook metadata layout")

    width = max(len(row) for row in metadata_rows)
    padded = [list(row) + [None] * (width - len(row)) for row in metadata_rows]
    proxy_rows: list[dict[str, object]] = []
    for column in range(1, width):
        identity = padded[0][column]
        if identity is None:
            continue
        proxy_rows.append(
            {
                "external_column_1based": column + 1,
                "identity": identity,
                "metadata_row2": padded[1][column],
                "metadata_row3": padded[2][column],
                "metadata_row4": padded[3][column],
                "latitude": as_float(padded[4][column]),
                "longitude": as_float(padded[5][column]),
                "elevation": as_float(padded[6][column]),
                "metadata_row8": padded[7][column],
                "first_year": as_int(padded[8][column]),
                "last_year": as_int(padded[9][column]),
                "metadata_row11": padded[10][column],
                "proxy_type": padded[11][column],
                "metadata_row13": padded[12][column],
                "metadata_row14": padded[13][column],
            }
        )

    with PROJECT_LEDGER.open(newline="", encoding="utf-8-sig") as handle:
        project = list(csv.DictReader(handle))
    identity_fields = (
        "record_id", "data_name", "site_name", "public_identity_id",
        "formal_metadata_identity", "metadata_identities",
    )
    project_tokens: dict[str, set[int]] = {}
    for index, row in enumerate(project):
        for field in identity_fields:
            token = norm(row.get(field))
            if token:
                project_tokens.setdefault(token, set()).add(index)

    overlap_rows: list[dict[str, object]] = []
    for ext in proxy_rows:
        ext_tokens = {
            norm(ext[key]) for key in ("identity", "metadata_row2", "metadata_row3", "metadata_row4")
            if norm(ext[key])
        }
        exact_indices: set[int] = set()
        for token in ext_tokens:
            exact_indices.update(project_tokens.get(token, set()))
        coordinate_indices: set[int] = set()
        lat = ext["latitude"]
        lon = ext["longitude"]
        if lat is not None and lon is not None:
            for index, row in enumerate(project):
                plat = as_float(row.get("latitude"))
                plon = as_float(row.get("longitude"))
                if plat is not None and plon is not None and abs(plat - lat) <= 0.01 and abs(plon - lon) <= 0.01:
                    coordinate_indices.add(index)
        all_indices = sorted(exact_indices | coordinate_indices)
        if not all_indices:
            overlap_rows.append(
                {
                    "external_identity": ext["identity"],
                    "match_type": "none",
                    "project_network": "",
                    "project_record": "",
                    "project_identity": "",
                    "coordinate_distance_degree": "",
                }
            )
            continue
        for index in all_indices:
            row = project[index]
            plat = as_float(row.get("latitude"))
            plon = as_float(row.get("longitude"))
            distance = None
            if lat is not None and lon is not None and plat is not None and plon is not None:
                distance = ((plat - lat) ** 2 + (plon - lon) ** 2) ** 0.5
            overlap_rows.append(
                {
                    "external_identity": ext["identity"],
                    "match_type": (
                        "exact_identity_and_coordinate" if index in exact_indices and index in coordinate_indices
                        else "exact_identity" if index in exact_indices else "coordinate_only"
                    ),
                    "project_network": row.get("network", ""),
                    "project_record": row.get("record_id", ""),
                    "project_identity": row.get("public_identity_id") or row.get("formal_metadata_identity") or row.get("data_name"),
                    "coordinate_distance_degree": distance,
                }
            )

    with (OUT / "external_proxy_metadata.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(proxy_rows[0]))
        writer.writeheader()
        writer.writerows(proxy_rows)
    with (OUT / "project_proxy_metadata_overlap.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(overlap_rows[0]))
        writer.writeheader()
        writer.writerows(overlap_rows)

    matched_external = {
        str(row["external_identity"]) for row in overlap_rows if row["match_type"] != "none"
    }
    exact_external = {
        str(row["external_identity"]) for row in overlap_rows if str(row["match_type"]).startswith("exact_identity")
    }
    hma_long = [
        row for row in proxy_rows
        if row["latitude"] is not None and row["longitude"] is not None
        and 25 <= float(row["latitude"]) <= 40 and 70 <= float(row["longitude"]) <= 105
        and row["first_year"] is not None and int(row["first_year"]) <= 1500
        and row["last_year"] is not None and int(row["last_year"]) >= 1940
    ]
    summary = {
        "round": 199,
        "record": "10.5281/zenodo.18477496",
        "targeted_member": MEMBER,
        "http_bytes_read": len(local) + len(header_tail) + len(compressed),
        "workbook_size": len(workbook_bytes),
        "workbook_sha256": workbook_sha256,
        "sheet_names": workbook.sheetnames,
        "decoded_rows": "1-14 only",
        "annual_rows_15_onward_decoded": False,
        "external_proxy_count": len(proxy_rows),
        "external_with_any_project_metadata_match": len(matched_external),
        "external_with_exact_identity_match": len(exact_external),
        "hma_long_candidate_count": len(hma_long),
        "locked_period_values_accessed": False,
        "jionglacuo_accessed": False,
        "decision": (
            "source_overlap_present_requires_exact_lineage_adjudication"
            if matched_external else "no_metadata_overlap_detected"
        ),
    }
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (OUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "remote_archive": {"record": "10.5281/zenodo.18477496", "member": MEMBER},
                "workbook_path": str(workbook_path),
                "workbook_sha256": workbook_sha256,
                "project_ledger_path": str(PROJECT_LEDGER),
                "project_ledger_sha256": hashlib.sha256(PROJECT_LEDGER.read_bytes()).hexdigest(),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    report = [
        "# Round 199 East Asia proxy identity metadata audit",
        "",
        f"- Decision: `{summary['decision']}`.",
        f"- External proxies: `{len(proxy_rows)}`.",
        f"- Any project metadata match: `{len(matched_external)}`.",
        f"- Exact normalized identity match: `{len(exact_external)}`.",
        f"- HMA records spanning 1500 and 1940: `{len(hma_long)}`.",
        "- Only workbook rows 1-14 were decoded; annual rows 15 onward were not iterated.",
        "- No climate field, locked-period value, or Jionglacuo value was accessed.",
    ]
    (OUT / "FINAL_RESULT.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
