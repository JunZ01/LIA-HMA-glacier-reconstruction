"""Safe coverage-only preflight for Jiang Co alkenone summer temperature."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
SOURCE = (
    ROOT
    / "validation_data/sources/external_hma_temperature_candidates_lake_ice_v1"
    / "jiang_co_S2095927323002943_mmc4.xlsx"
)
OUTPUT = (
    ROOT
    / "validation_data/results/paleo_surface_validation_v0"
    / "climate_reconstruction_temperature_jiang_safe_preflight_v2"
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contract = OUTPUT / "ROUND_CONTRACT.md"
    before = {"source": sha256(SOURCE), "contract": sha256(contract)}
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    ages = []
    skipped_locked = 0
    for row in sheet.iter_rows(min_row=12):
        try:
            age = float(row[1].value)
        except (TypeError, ValueError):
            continue
        rounded = int(round(age))
        if 1921 <= rounded <= 1939 and abs(age - rounded) <= 0.51:
            skipped_locked += 1
            continue
        value = row[6].value
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        ages.append(age)
    workbook.close()
    ages = np.sort(np.asarray(ages, dtype=float))
    if len(ages) == 0:
        raise ValueError("No finite Jiang temperature rows")
    spacings = np.diff(ages)
    windows = []
    for start in range(1500, 1900, 50):
        windows.append(
            {
                "window_start": start,
                "window_end": start + 49,
                "finite_temperature_count": int(
                    np.sum((ages >= start) & (ages <= start + 49))
                ),
            }
        )
    result = {
        "status": "complete",
        "age_column": 2,
        "temperature_column": 7,
        "age_min_ad_excluding_locked": float(ages.min()),
        "age_max_ad_excluding_locked": float(ages.max()),
        "finite_temperature_count_excluding_locked": int(len(ages)),
        "median_age_spacing_years": float(np.median(spacings)),
        "maximum_age_spacing_years": float(np.max(spacings)),
        "finite_temperature_count_1940_2000": int(np.sum((ages >= 1940) & (ages <= 2000))),
        "window_counts": windows,
        "locked_age_rows_skipped_before_temperature_access": skipped_locked,
        "temperature_values_reported": False,
        "scientific_score_calculated": False,
    }
    after = {"source": sha256(SOURCE), "contract": sha256(contract)}
    result["source_hashes_unchanged"] = before == after
    (OUTPUT / "preflight_summary.json").write_text(
        json.dumps(result, indent=2), encoding="utf-8"
    )
    (OUTPUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "source_path": str(SOURCE),
                "source_sha256_before": before["source"],
                "source_sha256_after": after["source"],
                "contract_sha256": before["contract"],
                "all_input_hashes_unchanged": before == after,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
