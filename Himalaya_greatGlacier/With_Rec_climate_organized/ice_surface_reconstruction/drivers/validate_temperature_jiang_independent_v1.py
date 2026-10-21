"""Blind central-HMA low-frequency validation against Jiang Co UK37 temperature."""

from __future__ import annotations

import json
from pathlib import Path

import h5py
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

import validate_temperature_external_low_frequency_v1 as base


SOURCE = (
    base.ROOT
    / "validation_data/sources/external_hma_temperature_candidates_lake_ice_v1"
    / "jiang_co_S2095927323002943_mmc4.xlsx"
)
PUBLIC = Path(
    r"E:\3spyderProject\HMA_500\Down_scale_new\out_data\climate_anomaly_05deg_with_ensemble_1000-2000.nc"
)
OUTPUT = (
    base.RESULTS
    / "climate_reconstruction_temperature_jiang_independent_blind_validation_v1"
)
YEARS = np.arange(1000, 2001)
VARIABLE = "tas_ensemble_mean_rec_result"
WINDOWS = [(year, year + 49) for year in range(1550, 1900, 50)]


def read_jiang(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = []
    for row in sheet.iter_rows(min_row=12):
        try:
            age = float(row[1].value)
        except (TypeError, ValueError):
            continue
        if 1921 <= age <= 1939.999:
            continue
        value = row[6].value
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        rows.append({"age": age, "external": float(value)})
    workbook.close()
    return pd.DataFrame(rows).sort_values("age").reset_index(drop=True)


def field_without_locked(dataset: h5py.Dataset) -> np.ndarray:
    result = np.full(dataset.shape, np.nan, dtype=float)
    result[:921] = np.asarray(dataset[:921], dtype=float)
    result[940:] = np.asarray(dataset[940:], dtype=float)
    return result


def local_series(field: np.ndarray, lat: np.ndarray, lon: np.ndarray) -> tuple[np.ndarray, int]:
    mask = (
        (lat[:, None] >= 31.0)
        & (lat[:, None] <= 32.0)
        & (lon[None, :] >= 90.25)
        & (lon[None, :] <= 91.25)
    )
    weights = np.cos(np.deg2rad(lat))[:, None] * mask
    values = np.where(mask[None, :, :], field, np.nan)
    valid = np.isfinite(values)
    numerator = np.nansum(values * weights[None, :, :], axis=(1, 2))
    denominator = np.sum(valid * weights[None, :, :], axis=(1, 2))
    return (
        np.divide(numerator, denominator, out=np.full(len(field), np.nan), where=denominator > 0),
        int(mask.sum()),
    )


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contract = OUTPUT / "ROUND_CONTRACT.md"
    paths = {
        "jiang": SOURCE,
        "public_field": PUBLIC,
        "frozen_candidate": base.CANDIDATE,
        "contract": contract,
    }
    before = {name: base.sha256(path) for name, path in paths.items()}
    external = read_jiang(SOURCE)

    with h5py.File(PUBLIC, "r") as handle:
        years = np.asarray(handle["time"], dtype=int).ravel()
        lat = np.asarray(handle["lat"], dtype=float).ravel()
        lon = np.asarray(handle["lon"], dtype=float).ravel()
        if not np.array_equal(years, YEARS):
            raise ValueError("Unexpected field years")
        project, cell_count = local_series(field_without_locked(handle[VARIABLE]), lat, lon)

    frozen = base.read_candidate(base.CANDIDATE)
    increment = frozen.candidate.to_numpy(float) - frozen.project.to_numpy(float)
    candidate = project + increment
    allowed = external.loc[
        ~external.age.between(base.LOCKED_START, base.LOCKED_END)
        & external.age.between(1000, 2000)
    ].copy()
    allowed["project"] = np.interp(allowed.age, YEARS, project)
    allowed["candidate"] = np.interp(allowed.age, YEARS, candidate)

    modern = allowed.age.between(1940, 2000)
    if int(modern.sum()) != 8:
        raise ValueError("Expected eight Jiang modern alignment points")
    for column in ["project", "candidate", "external"]:
        allowed[f"{column}_aligned"] = allowed[column] - allowed.loc[modern, column].mean()

    rows = []
    for start, end in WINDOWS:
        subset = allowed.loc[allowed.age.between(start, end)]
        if len(subset) < 5:
            raise ValueError(f"Jiang window {start}-{end} has only {len(subset)} points")
        rows.append(
            {
                "window_start": start,
                "window_end": end,
                "sample_count": len(subset),
                "project_degC": subset.project_aligned.mean(),
                "candidate_degC": subset.candidate_aligned.mean(),
                "jiang_degC": subset.external_aligned.mean(),
            }
        )
    windows = pd.DataFrame(rows)
    project_error = windows.project_degC - windows.jiang_degC
    candidate_error = windows.candidate_degC - windows.jiang_degC
    project_rmse = float(np.sqrt(np.mean(project_error**2)))
    candidate_rmse = float(np.sqrt(np.mean(candidate_error**2)))
    improvement = (project_rmse - candidate_rmse) / project_rmse
    project_bias = float(project_error.mean())
    candidate_bias = float(candidate_error.mean())
    project_corr = float(windows[["project_degC", "jiang_degC"]].corr().iloc[0, 1])
    candidate_corr = float(windows[["candidate_degC", "jiang_degC"]].corr().iloc[0, 1])
    gates = {
        "rmse_improves_at_least_10pct": bool(improvement >= 0.10 - 1e-12),
        "absolute_mean_bias_decreases": bool(abs(candidate_bias) < abs(project_bias)),
        "window_correlation_decline_at_most_0p10": bool(candidate_corr - project_corr >= -0.10 - 1e-12),
    }
    passes = bool(all(gates.values()))
    windows.to_csv(OUTPUT / "jiang_50year_windows.csv", index=False)
    pd.DataFrame(
        [{
            "project_rmse_degC": project_rmse,
            "candidate_rmse_degC": candidate_rmse,
            "rmse_improvement_fraction": improvement,
            "project_mean_bias_degC": project_bias,
            "candidate_mean_bias_degC": candidate_bias,
            "project_window_correlation": project_corr,
            "candidate_window_correlation": candidate_corr,
            "grid_cell_count": cell_count,
        }]
    ).to_csv(OUTPUT / "jiang_metrics.csv", index=False)

    x = (windows.window_start + windows.window_end) / 2
    fig, axis = plt.subplots(figsize=(10.5, 5.2), constrained_layout=True)
    axis.plot(x, windows.jiang_degC, color="#111111", marker="o", label="Jiang Co summer")
    axis.plot(x, windows.project_degC, color="#d55e00", marker="s", label="Project local annual")
    axis.plot(x, windows.candidate_degC, color="#0072b2", marker="^", label="Frozen increment")
    axis.axhline(0, color="#888888", linewidth=0.8)
    axis.set_title("Central HMA low-frequency temperature: Jiang Co")
    axis.set_xlabel("50-year window midpoint")
    axis.set_ylabel("Temperature relative to 1940-2000 sample mean (degC)")
    axis.grid(alpha=0.2)
    axis.legend()
    fig.savefig(OUTPUT / "jiang_blind_validation.png", dpi=220)
    plt.close(fig)

    after = {name: base.sha256(path) for name, path in paths.items()}
    (OUTPUT / "source_manifest.json").write_text(
        json.dumps({
            "inputs": {
                name: {"path": str(paths[name]), "sha256_before": before[name], "sha256_after": after[name]}
                for name in paths
            },
            "all_input_hashes_unchanged": before == after,
        }, indent=2), encoding="utf-8"
    )
    result = {
        "status": "complete",
        "project_rmse_degC": project_rmse,
        "candidate_rmse_degC": candidate_rmse,
        "rmse_improvement_fraction": improvement,
        "project_mean_bias_degC": project_bias,
        "candidate_mean_bias_degC": candidate_bias,
        "project_window_correlation": project_corr,
        "candidate_window_correlation": candidate_corr,
        "gates": gates,
        "passes_jiang_gates": passes,
        "locked_age_rows_skipped_before_temperature_access": True,
        "locked_1921_1939_field_slices_read": False,
        "phase_1901_1920_used": False,
        "annual_air_temperature_amplitude_validated": False,
        "spatial_field_release_authorized": False,
        "production_release_authorized": False,
    }
    (OUTPUT / "summary.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    (OUTPUT / "FINAL_RESULT.md").write_text(
        "# Jiang Co Independent Blind Validation\n\n"
        f"- RMSE improvement: `{100*improvement:.2f}%`.\n"
        f"- Mean bias: `{project_bias:+.3f}` to `{candidate_bias:+.3f} degC`.\n"
        f"- All gates pass: `{str(passes).lower()}`.\n"
        "- Central summer-pattern evidence only; no annual/spatial field is released.\n",
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
