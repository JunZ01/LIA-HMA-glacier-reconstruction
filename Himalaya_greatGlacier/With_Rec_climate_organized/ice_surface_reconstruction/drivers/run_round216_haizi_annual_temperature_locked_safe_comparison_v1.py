from __future__ import annotations

import importlib.util
import json
import math
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
LOOP = (
    ROOT
    / "validation_data/results/paleo_surface_validation_v0"
    / "climate_reconstruction_improvement_loop_v2"
)
OUT = LOOP / "round_216_haizi_annual_temperature_locked_safe_comparison_v1"
ROUND214 = LOOP / "round_214_haizi_annual_temperature_data_recovery_v1"
ROUND215 = LOOP / "round_215_haizi_daocheng_sequence_identity_v3"
WORKBOOK = ROUND214 / "41598_2024_79096_MOESM1_ESM.xlsx"
CONTRACT = OUT / "ROUND_CONTRACT.md"
DRIVER = Path(__file__).resolve()
ROUND212_DRIVER = DRIVER.parent / "run_round212_wscp_temperature_locked_safe_comparison_v1.py"
EXPECTED_WORKBOOK_SHA256 = "35302c0ba1912b015ee8d8b1a5c05df9aa82dd44d4467f7b07278e723611d3eb"
SITE_LAT = 29.58
SITE_LON = 100.65
LOCKED_START = 1921
LOCKED_END = 1939


def load_round212():
    spec = importlib.util.spec_from_file_location("round212_for_round216", ROUND212_DRIVER)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot load Round 212 helpers")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_haizi_year_first(helper) -> tuple[dict[int, float], int]:
    if helper.sha256(WORKBOOK) != EXPECTED_WORKBOOK_SHA256:
        raise RuntimeError("Haizi workbook hash changed")
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if workbook.sheetnames != ["Sheet1", "Sheet2", "Sheet3"]:
        raise RuntimeError("Unexpected Haizi workbook sheets")
    sheet = workbook["Sheet1"]
    if sheet.max_row != 1234 or sheet.max_column != 9:
        raise RuntimeError("Unexpected Haizi workbook dimensions")
    header = [sheet.cell(1, column).value for column in range(1, 5)]
    if (
        header[0] != "year"
        or header[1] != "std"
        or not isinstance(header[2], str)
        or not header[2].startswith("Temperature(")
        or header[3] != "num"
    ):
        raise RuntimeError("Unexpected Haizi workbook header")

    values: dict[int, float] = {}
    skipped = 0
    for row in sheet.iter_rows(min_row=2, max_col=4, values_only=False):
        year_token = row[0].value
        if year_token is None:
            continue
        year = int(year_token)
        if year < 1115 or year > 2017:
            continue
        if LOCKED_START <= year <= LOCKED_END:
            skipped += 1
            continue
        value = float(row[2].value)
        if not math.isfinite(value):
            raise RuntimeError(f"Non-finite Haizi temperature in allowed year {year}")
        values[year] = value
    workbook.close()
    expected = set(range(1115, 1921)) | set(range(1940, 2018))
    if set(values) != expected or skipped != 19:
        raise RuntimeError("Haizi allowed coverage mismatch")
    return values, skipped


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    helper = load_round212()
    required = [
        WORKBOOK,
        ROUND214 / "summary.json",
        ROUND215 / "summary.json",
        CONTRACT,
        ROUND212_DRIVER,
    ]
    for path in required:
        if not path.exists():
            raise FileNotFoundError(path)
    round215 = json.loads((ROUND215 / "summary.json").read_text(encoding="utf-8"))
    if not round215["admissible_as_independent_annual_temperature_evidence"]:
        raise RuntimeError("Round 215 did not admit Haizi as independent evidence")

    round206 = helper.load_round206()
    required.extend([round206.LINEAGE, round206.REFERENCE, round206.HMA_MASK])
    sources_before = {str(path): helper.state(path) for path in required}
    reconstruction, locked_rows_skipped = read_haizi_year_first(helper)

    years = np.arange(1000, 2001)
    locked = (years >= LOCKED_START) & (years <= LOCKED_END)
    reference, latitudes, longitudes, reference_hma_mask = round206.read_reference_allowed(
        round206.REFERENCE, years
    )
    hma_mask = np.load(round206.HMA_MASK).astype(bool)
    if not np.array_equal(hma_mask, reference_hma_mask):
        raise RuntimeError("HMA mask mismatch")

    lineage = pd.read_csv(round206.LINEAGE)
    members = lineage[lineage.variable == "tas"].copy()
    expected_members = {"AO_6", "AO_9", "HGAO_6", "HGAO_9", "O_6", "O_9"}
    if len(members) != 6 or set(members.member) != expected_members:
        raise RuntimeError("Temperature member lineage changed")
    round23 = round206.load_round23_module()
    project_members = []
    member_sources: dict[str, Path] = {}
    for row in members.itertuples(index=False):
        path = Path(row.source_tif)
        if helper.sha256(path) != row.source_tif_sha256:
            raise RuntimeError(f"Member hash mismatch: {row.member}")
        field, _ = round23.read_allowed_tif(path, years, latitudes, longitudes)
        project_members.append(field)
        member_sources[row.member] = path
    project = round23.finite_mean(project_members)
    if np.any(np.isfinite(project[locked])) or np.any(np.isfinite(reference[locked])):
        raise RuntimeError("Locked project/reference values became visible")

    box_mask = (
        (latitudes[:, None] >= SITE_LAT - 1.0)
        & (latitudes[:, None] <= SITE_LAT + 1.0)
        & (longitudes[None, :] >= SITE_LON - 1.0)
        & (longitudes[None, :] <= SITE_LON + 1.0)
    )
    if int(box_mask.sum()) < 4:
        raise RuntimeError("Source-centered box has too few cells")
    lat_index = int(np.argmin(np.abs(latitudes - SITE_LAT)))
    lon_index = int(np.argmin(np.abs(longitudes - SITE_LON)))

    external = helper.anomaly(helper.vector_from_dict(reconstruction, years), years, 1961, 1990)
    project_box = helper.anomaly(
        round206.area_weighted_series(project, box_mask, latitudes), years, 1961, 1990
    )
    reference_box = helper.anomaly(
        round206.area_weighted_series(reference, box_mask, latitudes), years, 1961, 1990
    )
    project_cell = helper.anomaly(project[:, lat_index, lon_index], years, 1961, 1990)
    reference_cell = helper.anomaly(reference[:, lat_index, lon_index], years, 1961, 1990)
    curves = {
        "haizi_annual_reconstruction_anomaly_degC": external,
        "project_local_box_annual_anomaly_degC": project_box,
        "reference_local_box_annual_anomaly_degC": reference_box,
        "project_nearest_cell_annual_anomaly_degC": project_cell,
        "reference_nearest_cell_annual_anomaly_degC": reference_cell,
    }
    smooth = {name: round206.segment_smooth(values, years) for name, values in curves.items()}

    annual_rows = []
    for index, year in enumerate(years):
        if LOCKED_START <= year <= LOCKED_END:
            continue
        row: dict[str, float | int] = {"year": int(year)}
        for name, values in curves.items():
            row[name] = values[index]
            row[f"{name}_smooth31"] = smooth[name][index]
        annual_rows.append(row)
    pd.DataFrame(annual_rows).to_csv(OUT / "annual_allowed_curves.csv", index=False)

    historical_names = list(curves)
    window_rows = []
    for start in range(1500, 1900, 50):
        end = start + 49
        selection = (years >= start) & (years <= end)
        row: dict[str, float | int] = {"start_year": start, "end_year": end, "center_year": start + 24.5}
        for name in historical_names:
            if int(np.isfinite(curves[name][selection]).sum()) != 50:
                raise RuntimeError(f"Incomplete window {name} {start}-{end}")
            row[name] = float(np.mean(curves[name][selection]))
        window_rows.append(row)
    windows = pd.DataFrame(window_rows)
    windows.to_csv(OUT / "temperature_50year_windows_1500_1899.csv", index=False)

    interval_rows = []
    for start, end in ((1850, 1900), (1901, 1920), (1940, 2000)):
        row = {"start_year": start, "end_year": end}
        for name in historical_names:
            row[f"{name}_trend_degC_per_century"] = helper.trend(curves[name], years, start, end)
        interval_rows.append(row)
    pd.DataFrame(interval_rows).to_csv(OUT / "diagnostic_interval_trends.csv", index=False)

    external_windows = windows[historical_names[0]].to_numpy(float)
    project_windows = windows[historical_names[1]].to_numpy(float)
    reference_windows = windows[historical_names[2]].to_numpy(float)
    modern_1940 = (years >= 1940) & (years <= 2000)
    calibration_overlap = (years >= 1959) & (years <= 2000)
    metrics = {
        "external_project_correlation_1940_2000": helper.correlation(external[modern_1940], project_box[modern_1940]),
        "external_reference_correlation_1940_2000": helper.correlation(external[modern_1940], reference_box[modern_1940]),
        "external_project_correlation_1959_2000": helper.correlation(external[calibration_overlap], project_box[calibration_overlap]),
        "external_reference_correlation_1959_2000": helper.correlation(external[calibration_overlap], reference_box[calibration_overlap]),
        "external_project_window_correlation": helper.correlation(external_windows, project_windows),
        "external_reference_window_correlation": helper.correlation(external_windows, reference_windows),
        "external_negative_window_count": int(np.sum(external_windows < 0)),
        "project_box_negative_window_count": int(np.sum(project_windows < 0)),
        "reference_box_negative_window_count": int(np.sum(reference_windows < 0)),
        "median_external_anomaly_1500_1899_degC": float(np.median(external_windows)),
        "median_project_box_anomaly_1500_1899_degC": float(np.median(project_windows)),
        "median_reference_box_anomaly_1500_1899_degC": float(np.median(reference_windows)),
        "median_project_minus_external_1500_1899_degC": float(np.median(project_windows) - np.median(external_windows)),
        "external_project_adjacent_window_sign_agreement": int(np.sum(np.sign(np.diff(external_windows)) == np.sign(np.diff(project_windows)))),
        "external_reference_adjacent_window_sign_agreement": int(np.sum(np.sign(np.diff(external_windows)) == np.sign(np.diff(reference_windows)))),
    }
    supports_local_warm_bias = (
        metrics["external_negative_window_count"] >= 6
        and metrics["external_negative_window_count"] - metrics["project_box_negative_window_count"] >= 3
        and metrics["median_project_minus_external_1500_1899_degC"] >= 0.15
    )
    decision = (
        "independent_haizi_annual_record_supports_local_project_warm_bias_direction"
        if supports_local_warm_bias
        else "independent_haizi_annual_record_does_not_support_frozen_local_warm_bias_gate"
    )

    figure, axes = plt.subplots(2, 1, figsize=(12.4, 8.0), sharex=True, gridspec_kw={"height_ratios": [2.2, 1.0], "hspace": 0.16})
    colors = ["#222222", "#B33A3A", "#147D78"]
    labels = ["Haizi annual tree-ring", "Project local annual", "Four-family local annual"]
    for name, color, label in zip(historical_names[:3], colors, labels):
        axes[0].plot(years, smooth[name], color=color, linewidth=2.2, label=label)
    axes[0].axhline(0, color="#777777", linewidth=0.8)
    axes[0].axvspan(1921, 1939, color="#E5E5E5", alpha=0.85, linewidth=0)
    axes[0].set_ylabel("31-year smooth anomaly (degC; 1961-1990)")
    axes[0].set_title("Independent Haizi annual-temperature comparison")
    axes[0].grid(axis="y", color="#DDDDDD", linewidth=0.6)
    axes[0].legend(frameon=False, ncol=3, loc="upper left")
    centers = windows.center_year.to_numpy(float)
    for values, color, label, marker in zip([external_windows, project_windows, reference_windows], colors, labels, ["o", "s", "^"]):
        axes[1].plot(centers, values, color=color, marker=marker, linewidth=1.9, label=label)
    axes[1].axhline(0, color="#777777", linewidth=0.8)
    axes[1].set_xlim(1490, 2005)
    axes[1].set_xlabel("Year")
    axes[1].set_ylabel("50-year anomaly (degC)")
    axes[1].grid(axis="y", color="#DDDDDD", linewidth=0.6)
    axes[1].legend(frameon=False, ncol=3, loc="lower left")
    figure.text(0.99, 0.01, "Gray gap: sealed 1921-1939. Haizi target: previous October-current September mean.", ha="right", va="bottom", fontsize=8.5, color="#555555")
    figure.subplots_adjust(left=0.09, right=0.98, top=0.93, bottom=0.1)
    figure.savefig(OUT / "haizi_project_reference_curves_1500_2000.png", dpi=220)
    plt.close(figure)

    all_sources = {str(path): path for path in required}
    all_sources.update({f"member_{name}": path for name, path in member_sources.items()})
    sources_after = {str(path): helper.state(path) for path in required}
    if sources_before != sources_after:
        raise RuntimeError("A read-only source changed")
    member_states = {name: helper.state(path) for name, path in member_sources.items()}
    summary = {
        "round": 216,
        "decision": decision,
        "source_independence": round215["decision"],
        "target_boundary": {
            "external": "previous-October to current-September mean temperature",
            "project_reference": "calendar annual mean temperature",
            "hma_wide_correction_permitted": False,
        },
        "spatial_support": {
            "centroid_latitude": SITE_LAT,
            "centroid_longitude": SITE_LON,
            "box_cell_count": int(box_mask.sum()),
            "nearest_grid_latitude": float(latitudes[lat_index]),
            "nearest_grid_longitude": float(longitudes[lon_index]),
            "member_count": len(project_members),
        },
        "locked_parser": {
            "locked_rows_skipped_before_temperature_cell_access": locked_rows_skipped,
            "project_locked_raster_bands_read": False,
            "smoothing_across_locked_gap": False,
        },
        "metrics": metrics,
        "source_states_unchanged": True,
        "locked_period_values_accessed": False,
        "jionglacuo_accessed": False,
        "production_changed": False,
    }
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    (OUT / "source_manifest.json").write_text(json.dumps({"sources_before": sources_before, "sources_after": sources_after, "member_sources": member_states}, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    result = f"""# Round 216 Result

## Decision

`{decision}`

Across the eight 50-year windows in 1500-1899, the independent Haizi annual
record is below its 1961-1990 baseline in `{metrics['external_negative_window_count']}/8`
windows, versus `{metrics['project_box_negative_window_count']}/8` for the project
local annual curve and `{metrics['reference_box_negative_window_count']}/8` for
the four-family reference. Median anomalies are
`{metrics['median_external_anomaly_1500_1899_degC']:+.3f}`,
`{metrics['median_project_box_anomaly_1500_1899_degC']:+.3f}`, and
`{metrics['median_reference_box_anomaly_1500_1899_degC']:+.3f}` degC.

The project-minus-Haizi historical median difference is
`{metrics['median_project_minus_external_1500_1899_degC']:+.3f}` degC. This
local annual record does not by itself authorize an HMA-wide correction. No
production file changed and no locked-period or Jionglacuo value was accessed.
"""
    (OUT / "FINAL_RESULT.md").write_text(result, encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
