from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from compare_noaa_hma_lake_temperature_project_v1 import LINEAGE, PROJECT_MEMBERS, read_project_series


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
BASE = ROOT / "validation_data/results/paleo_surface_validation_v0/climate_reconstruction_improvement_loop_v2"
SOURCE_ROOT = BASE / "temperature_lake_brgdgt_recalibration_sensitivity_v1"
SOURCE = SOURCE_ROOT / "source/Liang_et_al_20241108_CEE_Supplementary_Data.xlsx"
OUTPUT = BASE / "temperature_lake_brgdgt_eight_record_safe_audit_v3"
DRIVER = Path(__file__).resolve()
LOCKED_START = 1921.0
LOCKED_END = 1939.999999

RECORDS = {
    "Bangong": {"age": "G", "raw": "H", "constrained": "M", "lat": 33.73, "lon": 78.96, "target": "MAF"},
    "Ngamring": {"age": "V", "raw": "W", "constrained": "AB", "lat": 29.31, "lon": 87.18, "target": "MAF"},
    "Gahai": {"age": "AK", "raw": "AL", "constrained": None, "lat": 34.067, "lon": 102.325, "target": "ice_free_season"},
    "Tiancai": {"age": "AX", "raw": "AY", "constrained": None, "lat": 26.63, "lon": 99.72, "target": "MAF"},
    "Cuoqia": {"age": "BM", "raw": "BN", "constrained": None, "lat": 27.4052, "lon": 99.7722, "target": "MAF"},
    "Qinghai": {"age": "BZ", "raw": "CA", "constrained": None, "lat": 36.8112, "lon": 100.137, "target": "MAF"},
    "Tengchongqinghai": {"age": "CN", "raw": "CO", "constrained": None, "lat": 25.132, "lon": 98.57, "target": "MAF_or_MAAT"},
    "Lugu": {"age": "DB", "raw": "DC", "constrained": None, "lat": 27.7133, "lon": 100.79, "target": "MAF_or_MAAT"},
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def direction(value: float | None) -> str:
    if value is None:
        return "insufficient_points"
    if value > 1e-12:
        return "warming"
    if value < -1e-12:
        return "cooling"
    return "neutral"


def parse_workbook() -> tuple[pd.DataFrame, dict[str, int]]:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Data S7.  TP Holocene records"]
    rows = []
    locked = {record: 0 for record in RECORDS}
    for record, definition in RECORDS.items():
        for row_number in range(3, sheet.max_row + 1):
            age_value = sheet[f"{definition['age']}{row_number}"].value
            if not isinstance(age_value, (int, float)) or not np.isfinite(float(age_value)):
                continue
            age_bp = float(age_value)
            year_ce = 1950.0 - age_bp
            if LOCKED_START <= year_ce <= LOCKED_END:
                locked[record] += 1
                continue
            if not (1500.0 <= year_ce <= 1920.999999 or 1940.0 <= year_ce <= 1950.0):
                continue
            raw_value = sheet[f"{definition['raw']}{row_number}"].value
            if not isinstance(raw_value, (int, float)) or not np.isfinite(float(raw_value)):
                continue
            constrained_value = None
            if definition["constrained"] is not None:
                candidate = sheet[f"{definition['constrained']}{row_number}"].value
                if isinstance(candidate, (int, float)) and np.isfinite(float(candidate)):
                    constrained_value = float(candidate)
            rows.append(
                {
                    "record": record,
                    "source_row": row_number,
                    "age_bp": age_bp,
                    "year_ce": year_ce,
                    "raw_rf_maf_c": float(raw_value),
                    "constrained_maf_c": constrained_value,
                    "admitted_maf_c": constrained_value if constrained_value is not None else float(raw_value),
                    "latitude": definition["lat"],
                    "longitude": definition["lon"],
                    "target": definition["target"],
                }
            )
    workbook.close()
    frame = pd.DataFrame(rows).sort_values(["record", "year_ce"]).reset_index(drop=True)
    return frame, locked


def slope_per_century(years: np.ndarray, values: np.ndarray) -> float | None:
    if len(years) < 2 or np.unique(years).size < 2:
        return None
    return float(np.polyfit(years, values, 1)[0] * 100.0)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=False)
    source_hash_before = sha256(SOURCE)
    points, locked = parse_workbook()
    points.to_csv(OUTPUT / "safe_eight_record_points.csv", index=False)
    lineage = pd.read_csv(LINEAGE)
    lineage = lineage[(lineage.variable == "tas") & lineage.member.isin(PROJECT_MEMBERS)].copy()

    metrics = []
    project_series: dict[str, dict[str, pd.DataFrame]] = {}
    for record, definition in RECORDS.items():
        project_series[record] = {
            row.member: read_project_series(Path(row.source_tif), definition["lat"], definition["lon"])
            for row in lineage.itertuples(index=False)
        }
        for label, start, end in (("full_safe", 1500.0, 1920.999999), ("bridge_1850_1920", 1850.0, 1920.999999)):
            selected = points[(points.record == record) & points.year_ce.between(start, end)].sort_values("year_ce")
            years = selected.year_ce.to_numpy(float)
            admitted = selected.admitted_maf_c.to_numpy(float)
            raw = selected.raw_rf_maf_c.to_numpy(float)
            admitted_slope = slope_per_century(years, admitted)
            raw_slope = slope_per_century(years, raw)
            constrained_available = int(selected.constrained_maf_c.notna().sum())
            member_slopes = []
            if len(years) >= 2:
                for member in PROJECT_MEMBERS:
                    series = project_series[record][member]
                    sampled = np.interp(years, series.year.to_numpy(float), series.temperature.to_numpy(float))
                    member_slopes.append(slope_per_century(years, sampled))
            project_slope = float(np.median(member_slopes)) if member_slopes else None
            metrics.append(
                {
                    "record": record,
                    "window": label,
                    "n": int(len(selected)),
                    "first_year": float(years.min()) if len(years) else None,
                    "last_year": float(years.max()) if len(years) else None,
                    "raw_rf_trend_c_per_century": raw_slope,
                    "admitted_trend_c_per_century": admitted_slope,
                    "constrained_point_count": constrained_available,
                    "project_median_trend_at_proxy_years_c_per_century": project_slope,
                    "admitted_direction": direction(admitted_slope),
                    "project_direction": direction(project_slope),
                    "direction_agreement": (
                        direction(admitted_slope) == direction(project_slope) if admitted_slope is not None else None
                    ),
                }
            )
    metrics_frame = pd.DataFrame(metrics)
    metrics_frame.to_csv(OUTPUT / "eight_record_window_metrics.csv", index=False)

    full = metrics_frame[(metrics_frame.window == "full_safe") & (metrics_frame.n >= 2)].copy()
    fig, ax = plt.subplots(figsize=(10.8, 6.2))
    y = np.arange(len(full))
    ax.scatter(full.admitted_trend_c_per_century, y + 0.1, color="#111111", s=55, label="Admitted lake series")
    ax.scatter(full.project_median_trend_at_proxy_years_c_per_century, y - 0.1, color="#bf3f36", s=55, label="Project median")
    ax.axvline(0, color="#777777", lw=0.8)
    ax.set_yticks(y, full.record)
    ax.set_xlabel("1500-1920 safe-window trend (degC per century)")
    ax.set_title("Eight-record common-calibration audit")
    ax.grid(axis="x", color="#e2e2e2", lw=0.6)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(OUTPUT / "eight_record_full_safe_trends.png", dpi=220)
    plt.close(fig)

    if source_hash_before != sha256(SOURCE):
        raise RuntimeError("source-workbook-changed")
    records_json = metrics_frame.astype(object).where(pd.notna(metrics_frame), None).to_dict(orient="records")
    summary = {
        "experiment": "temperature_lake_brgdgt_eight_record_safe_audit_v3",
        "status": "complete",
        "records": records_json,
        "method": (
            "one global RF MAF calibration for all records; bacterial-cluster constrained values admitted only where "
            "the workbook supplies them for Bangong and Ngamring"
        ),
        "independence": (
            "underlying sediment records are distinct site-level physical archives; the 2024 workbook is a secondary "
            "common-method compilation and is not counted as another independent vote"
        ),
        "target_limit": "MAF or warm/ice-free-season temperature is compared with annual project temperature by direction only",
        "locked_rows_skipped_before_temperature_decode": locked,
        "sealed_period_values_used": False,
        "production_modified": False,
        "enkf_run": False,
        "oggm_run": False,
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    (OUTPUT / "summary.json").write_text(json.dumps(summary, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    manifest = {
        "paper": "https://doi.org/10.1038/s43247-024-01925-3",
        "dataset": "https://doi.org/10.6084/m9.figshare.26521102.v1",
        "source": {"path": str(SOURCE), "sha256": source_hash_before},
        "member_lineage": {"path": str(LINEAGE), "sha256": sha256(LINEAGE)},
        "driver": {"path": str(DRIVER), "sha256": sha256(DRIVER)},
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    (OUTPUT / "source_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    lines = [
        "# Eight-record lake common-calibration audit",
        "",
        "Underlying records are distinct lake sediment archives, but the 2024 workbook is a secondary common-method "
        "reanalysis. It is used for calibration and regional sensitivity, not as an extra independent vote.",
        "",
        "| Record | Window | n | Admitted trend | Project trend | Agreement |",
        "|---|---|---:|---:|---:|---|",
    ]
    for row in metrics_frame.itertuples(index=False):
        if row.n < 2:
            lines.append(f"| {row.record} | {row.window} | {row.n} | insufficient | insufficient | not evaluated |")
        else:
            lines.append(
                f"| {row.record} | {row.window} | {row.n} | {row.admitted_trend_c_per_century:+.3f} | "
                f"{row.project_median_trend_at_proxy_years_c_per_century:+.3f} | "
                f"{'yes' if row.direction_agreement else 'no'} |"
            )
    lines.extend(
        [
            "",
            "The sealed period was skipped before temperature decoding. Production files were not changed; EnKF and "
            "OGGM were not run; 20CRv3 was not accessed.",
        ]
    )
    (OUTPUT / "FINAL_RESULT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
