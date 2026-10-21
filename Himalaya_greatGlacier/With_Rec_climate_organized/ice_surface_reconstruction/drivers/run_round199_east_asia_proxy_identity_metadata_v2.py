from __future__ import annotations

from openpyxl import load_workbook as real_load_workbook

import run_round199_east_asia_proxy_identity_metadata_v1 as v1


class FirstSheetOnlyWorkbook:
    def __init__(self, workbook):
        self._workbook = workbook
        if not workbook.sheetnames:
            raise RuntimeError("Proxy workbook contains no worksheets")
        self.sheetnames = [workbook.sheetnames[0]]

    def __getitem__(self, name):
        return self._workbook[name]

    def close(self):
        self._workbook.close()


def first_sheet_load_workbook(*args, **kwargs):
    return FirstSheetOnlyWorkbook(real_load_workbook(*args, **kwargs))


v1.OUT = v1.ROOT / (
    "validation_data/results/paleo_surface_validation_v0/"
    "climate_reconstruction_improvement_loop_v2/"
    "round_199_east_asia_proxy_identity_metadata_v2"
)
v1.load_workbook = first_sheet_load_workbook


if __name__ == "__main__":
    v1.main()
