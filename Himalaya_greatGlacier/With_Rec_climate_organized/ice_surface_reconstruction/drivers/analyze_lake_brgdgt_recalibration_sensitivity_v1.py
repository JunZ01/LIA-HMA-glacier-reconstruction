from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

matplotlib.use("Agg")
import matplotlib.pyplot as plt

from compare_noaa_hma_lake_temperature_project_v1 import (
    LINEAGE,
    PROJECT_MEMBERS,
    interpolate_series,
    read_project_series,
)


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
BASE = ROOT / "validation_data/results/paleo_surface_validation_v0/climate_reconstruction_improvement_loop_v2"
EXPERIMENT = BASE / "temperature_lake_brgdgt_recalibration_sensitivity_v1"
SOURCE = EXPERIMENT / "source/Liang_et_al_20241108_CEE_Supplementary_Data.xlsx"
SCHEMA = EXPERIMENT / "workbook_schema.json"
DRIVER = Path(__file__).resolve()
LOCKED_START = 1921.0
LOCKED_END = 1939.999999

RECORDS = {
    "Tiancai": {"age_column": "AX", "temperature_column": "AY", "latitude": 26.63, "longitude": 99.72},
    "Cuoqia": {"age_column": "BM", "temperature_column": "BN", "latitude": 27.4052, "longitude": 99.7722},
    "Qinghai": {"age_column": "BZ", "temperature_column": "CA", "latitude": 36.8112, "longitude": 100.137},
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sign(value: float, tolerance: float = 1e-12) -> str:
    if value > tolerance:
        return "warming"
    if value < -tolerance:
        return "cooling"
    return "neutral"


def parse_safe_records() -> tuple[pd.DataFrame, dict[str, int]]:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Data S7.  TP Holocene records"]
    rows: list[dict[str, float | str | int]] = []
    audit = {name: 0 for name in RECORDS}
    for name, definition in RECORDS.items():
        for row_number in range(3, sheet.max_row + 1):
            # Age is the gate. Temperature is not decoded until this gate passes.
            age_value = sheet[f"{definition['age_column']}{row_number}"].value
            if not isinstance(age_value, (int, float)) or not np.isfinite(float(age_value)):
                continue
            age_bp = float(age_value)
            year_ce = 1950.0 - age_bp
            if LOCKED_START <= year_ce <= LOCKED_END:
                audit[name] += 1
                continue
            if not (1500.0 <= year_ce <= 1920.999999 or 1940.0 <= year_ce <= 1950.0):
                continue
            temperature_value = sheet[f"{definition['temperature_column']}{row_number}"].value
            if not isinstance(temperature_value, (int, float)) or not np.isfinite(float(temperature_value)):
                continue
            rows.append(
                {
                    "record": name,
                    "source_row": row_number,
                    "age_bp": age_bp,
                    "year_ce": year_ce,
                    "maf_rf_c": float(temperature_value),
                    "latitude": definition["latitude"],
                    "longitude": definition["longitude"],
                }
            )
    workbook.close()
    frame = pd.DataFrame(rows).sort_values(["record", "year_ce"]).reset_index(drop=True)
    if frame.empty:
        raise RuntimeError("no-safe-recalibrated-records")
    return frame, audit


def linear_change(frame: pd.DataFrame) -> dict[str, float | int | str | None]:
    ordered = frame.sort_values("year_ce")
    years = ordered.year_ce.to_numpy(float)
    values = ordered.maf_rf_c.to_numpy(float)
    slope = float(np.polyfit(years, values, 1)[0]) if len(ordered) >= 2 else None
    return {
        "start_year_ce": float(years.min()),
        "end_year_ce": float(years.max()),
        "n": int(len(ordered)),
        "rf_linear_change_c_per_century": slope * 100.0 if slope is not None else None,
        "rf_direction": sign(slope) if slope is not None else "insufficient_points",
    }


def main() -> None:
    source_hash_before = sha256(SOURCE)
    points, locked_audit = parse_safe_records()
    points.to_csv(EXPERIMENT / "safe_recalibrated_maf_points.csv", index=False)

    lineage = pd.read_csv(LINEAGE)
    lineage = lineage[(lineage.variable == "tas") & lineage.member.isin(PROJECT_MEMBERS)].copy()
    if len(lineage) != len(PROJECT_MEMBERS):
        raise RuntimeError("project-member-lineage-incomplete")

    project: dict[str, dict[str, pd.DataFrame]] = {}
    result_rows = []
    for name, definition in RECORDS.items():
        project[name] = {}
        for source in lineage.sort_values("member").itertuples(index=False):
            project[name][source.member] = read_project_series(
                Path(source.source_tif), definition["latitude"], definition["longitude"]
            )
        record = points[(points.record == name) & (points.year_ce <= 1920.999999)].copy()
        metric = linear_change(record)
        start = float(metric["start_year_ce"])
        end = float(metric["end_year_ce"])
        interval = end - start
        project_changes = []
        if interval > 0:
            for member in PROJECT_MEMBERS:
                series = project[name][member]
                project_changes.append(interpolate_series(series, end) - interpolate_series(series, start))
        project_change = float(np.median(project_changes)) if project_changes else None
        project_direction = sign(project_change) if project_change is not None else "insufficient_points"
        result_rows.append(
            {
                "record": name,
                **metric,
                "project_six_member_median_endpoint_change_c": project_change,
                "project_six_member_median_endpoint_change_c_per_century": (
                    project_change / interval * 100.0 if project_change is not None else None
                ),
                "project_direction": project_direction,
                "direction_agreement": (
                    metric["rf_direction"] == project_direction if project_change is not None else None
                ),
                "comparison_role": "same-site standardized-RF calibration sensitivity; not an independent reconstruction vote",
            }
        )
    metrics = pd.DataFrame(result_rows)
    metrics.to_csv(EXPERIMENT / "recalibrated_maf_project_metrics.csv", index=False)

    fig, axes = plt.subplots(1, 3, figsize=(14.4, 4.8), sharey=False)
    for ax, name in zip(axes, RECORDS):
        record = points[(points.record == name) & (points.year_ce <= 1920.999999)].sort_values("year_ce")
        proxy_base = float(record.iloc[0].maf_rf_c)
        ax.plot(record.year_ce, record.maf_rf_c - proxy_base, color="#111111", marker="o", ms=3.5, lw=1.8, label="2024 RF MAF")
        member_series = []
        years = np.arange(int(np.ceil(record.year_ce.min())), 1921)
        for member in PROJECT_MEMBERS:
            series = project[name][member]
            values = np.interp(years, series.year.to_numpy(float), series.temperature.to_numpy(float))
            values = values - np.interp(record.year_ce.min(), series.year.to_numpy(float), series.temperature.to_numpy(float))
            member_series.append(values)
            ax.plot(years, values, color="#5f8f9d", alpha=0.18, lw=0.8)
        ax.plot(years, np.median(np.vstack(member_series), axis=0), color="#bf3f36", lw=2.0, label="Project median")
        ax.axhline(0, color="#888888", lw=0.7)
        ax.set_title(name)
        ax.set_xlabel("Year CE")
        ax.grid(color="#e3e3e3", lw=0.6)
    axes[0].set_ylabel("Change from first safe point (degC)")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=2, frameon=False)
    fig.suptitle("Standardized brGDGT RF calibration versus project temperature", y=0.98)
    fig.tight_layout(rect=(0, 0.08, 1, 0.94))
    fig.savefig(EXPERIMENT / "recalibrated_maf_project_curves.png", dpi=220)
    plt.close(fig)

    source_hash_after = sha256(SOURCE)
    if source_hash_before != source_hash_after:
        raise RuntimeError("downloaded-workbook-changed")
    summary = {
        "experiment": "temperature_lake_brgdgt_recalibration_sensitivity_v1",
        "status": "complete",
        "records": metrics.astype(object).where(pd.notna(metrics), None).to_dict(orient="records"),
        "method_identity": (
            "All three records were reanalyzed by Liang et al. (2024) with one global lacustrine RF MAF calibration. "
            "The paper applies a 2.7 C bacterial-cluster constraint only to Bangong and Ngamring, not to Tiancai, Cuoqia, or Qinghai."
        ),
        "evidence_role": "calibration/operator sensitivity; secondary compilation; not an additional independent vote",
        "locked_rows_skipped_before_temperature_decode": locked_audit,
        "sealed_period_values_used": False,
        "no_interpolation_across_locked_period": True,
        "production_modified": False,
        "enkf_run": False,
        "oggm_run": False,
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    (EXPERIMENT / "summary.json").write_text(
        json.dumps(summary, indent=2, allow_nan=False) + "\n", encoding="utf-8"
    )
    manifest = {
        "paper": "https://doi.org/10.1038/s43247-024-01925-3",
        "dataset": "https://doi.org/10.6084/m9.figshare.26521102.v1",
        "source": {"path": str(SOURCE), "sha256": source_hash_before, "size_bytes": SOURCE.stat().st_size},
        "schema": {"path": str(SCHEMA), "sha256": sha256(SCHEMA)},
        "member_lineage": {"path": str(LINEAGE), "sha256": sha256(LINEAGE)},
        "driver": {"path": str(DRIVER), "sha256": sha256(DRIVER)},
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    (EXPERIMENT / "source_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

    lines = [
        "# Lake brGDGT common-calibration sensitivity",
        "",
        "Liang et al. (2024) reprocessed the source brGDGT records with one global RF calibration. "
        "Tiancai, Cuoqia, and Qinghai were not given the 2.7 C bacterial-cluster correction, so this is a "
        "standardized calibration sensitivity rather than a corrected independent reconstruction.",
        "",
        "| Record | Safe span | RF trend (C/century) | Project endpoint change (C/century) | Direction agrees |",
        "|---|---:|---:|---:|---|",
    ]
    for row in metrics.itertuples(index=False):
        if pd.isna(row.rf_linear_change_c_per_century):
            lines.append(
                f"| {row.record} | {row.start_year_ce:.1f}-{row.end_year_ce:.1f} | insufficient | "
                "insufficient | not evaluated |"
            )
        else:
            lines.append(
                f"| {row.record} | {row.start_year_ce:.1f}-{row.end_year_ce:.1f} | "
                f"{row.rf_linear_change_c_per_century:+.3f} | "
                f"{row.project_six_member_median_endpoint_change_c_per_century:+.3f} | "
                f"{'yes' if row.direction_agreement else 'no'} |"
            )
    lines.extend(
        [
            "",
            "The sealed period was skipped before temperature decoding. No production file was modified, "
            "EnKF and OGGM were not run, and 20CRv3 was not accessed.",
        ]
    )
    (EXPERIMENT / "FINAL_RESULT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
