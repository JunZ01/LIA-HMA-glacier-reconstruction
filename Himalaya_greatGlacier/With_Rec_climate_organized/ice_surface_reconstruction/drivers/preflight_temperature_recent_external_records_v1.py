"""Schema-only preflight for recent public HMA temperature supplements."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
SOURCE = ROOT / "validation_data/sources/external_hma_temperature_candidates_round2_v1"
OUTPUT = (
    ROOT
    / "validation_data/results/paleo_surface_validation_v0"
    / "climate_reconstruction_temperature_recent_external_preflight_v1"
)
LOCKED = range(1921, 1940)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def year_from(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    try:
        result = int(float(value))
    except (TypeError, ValueError):
        return None
    return result if 500 <= result <= 2100 else None


def text_preview(values: list[object], limit: int = 12) -> list[str]:
    result = []
    for value in values:
        if isinstance(value, str) and value.strip():
            result.append(value.strip())
        if len(result) >= limit:
            break
    return result


def inspect_workbook(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        year_values = []
        skipped_locked_rows = 0
        first_text_rows = []
        nonempty_rows = 0
        for row_index, row in enumerate(sheet.iter_rows(), start=1):
            if not row:
                continue
            first = row[0].value
            first_year = year_from(first)
            if first_year in LOCKED:
                skipped_locked_rows += 1
                continue
            values = [cell.value for cell in row]
            if not any(value is not None for value in values):
                continue
            nonempty_rows += 1
            if first_year is not None:
                year_values.append(first_year)
            if len(first_text_rows) < 12:
                preview = text_preview(values)
                if preview:
                    first_text_rows.append({"row": row_index, "text": preview})
        sheets.append(
            {
                "sheet": sheet.title,
                "max_row": int(sheet.max_row),
                "max_column": int(sheet.max_column),
                "nonempty_rows_excluding_locked": nonempty_rows,
                "first_column_year_min_excluding_locked": min(year_values) if year_values else None,
                "first_column_year_max_excluding_locked": max(year_values) if year_values else None,
                "first_column_year_count_excluding_locked": len(year_values),
                "locked_rows_skipped_before_remaining_cell_access": skipped_locked_rows,
                "text_header_preview": first_text_rows,
            }
        )
    workbook.close()
    return {"file": path.name, "kind": "xlsx", "sheets": sheets}


def inspect_docx(path: Path) -> dict[str, object]:
    document = Document(path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables = []
    for index, table in enumerate(document.tables, start=1):
        header = []
        if table.rows:
            header = [cell.text.strip() for cell in table.rows[0].cells]
        tables.append(
            {
                "table": index,
                "row_count": len(table.rows),
                "column_count": len(table.columns),
                "header_text": header,
                "body_values_read": False,
            }
        )
    return {
        "file": path.name,
        "kind": "docx",
        "paragraph_count": len(paragraphs),
        "paragraph_text_preview": paragraphs[:40],
        "table_count": len(tables),
        "tables": tables,
        "table_body_values_read": False,
    }


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    contract = OUTPUT / "ROUND_CONTRACT.md"
    if not contract.exists():
        raise FileNotFoundError("Contract must exist before preflight")
    files = [
        SOURCE / "Temperature reconstruction.xlsx",
        SOURCE / "41598_2024_79096_MOESM1_ESM.xlsx",
        SOURCE / "1-s2.0-S0048969724075508-mmc1.docx",
        SOURCE / "source_manifest.json",
        contract,
    ]
    before = {path.name: sha256(path) for path in files}
    inventory = [inspect_workbook(files[0]), inspect_workbook(files[1]), inspect_docx(files[2])]
    after = {path.name: sha256(path) for path in files}
    result = {
        "status": "complete",
        "scope": "schema_and_identity_preflight_only",
        "inventory": inventory,
        "locked_1921_1939_workbook_rows_skipped_before_remaining_cell_access": True,
        "docx_table_body_values_read": False,
        "scientific_values_reported": False,
        "scientific_score_calculated": False,
        "source_hashes_unchanged": before == after,
    }
    (OUTPUT / "preflight_inventory.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=True), encoding="utf-8"
    )
    (OUTPUT / "source_manifest.json").write_text(
        json.dumps(
            {
                "inputs": {
                    path.name: {
                        "path": str(path),
                        "sha256_before": before[path.name],
                        "sha256_after": after[path.name],
                    }
                    for path in files
                },
                "all_input_hashes_unchanged": before == after,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
