from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"E:\3spyderProject\1_LIA_himalaya_oggm")
EXPERIMENT = ROOT / (
    "validation_data/results/paleo_surface_validation_v0/climate_reconstruction_improvement_loop_v2/"
    "temperature_original_lake_author_supplements_v1"
)
SOURCE = EXPERIMENT / "source"
OUTPUT = EXPERIMENT / "workbook_schema.json"
FILES = sorted(SOURCE.glob("*.xlsx"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    workbooks = []
    for path in FILES:
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheets = []
        for sheet in workbook.worksheets:
            strings = []
            max_row = int(sheet.max_row or 0)
            max_column = int(sheet.max_column or 0)
            for row in sheet.iter_rows(min_row=1, max_row=min(15, max_row)):
                for cell in row:
                    if cell.data_type in {"s", "inlineStr"} and cell.value is not None:
                        strings.append({"coordinate": cell.coordinate, "text": str(cell.value)})
            sheets.append(
                {
                    "title": sheet.title,
                    "max_row": max_row,
                    "max_column": max_column,
                    "string_cells_first_15_rows": strings,
                }
            )
        workbook.close()
        workbooks.append(
            {
                "path": str(path),
                "sha256": sha256(path),
                "size_bytes": path.stat().st_size,
                "sheets": sheets,
            }
        )
    result = {
        "experiment": "temperature_original_lake_author_supplements_v1",
        "workbooks": workbooks,
        "numeric_cells_decoded": False,
        "sealed_period_values_used": False,
        "prohibited_source": "20CRv3",
        "prohibited_source_accessed": False,
    }
    OUTPUT.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
